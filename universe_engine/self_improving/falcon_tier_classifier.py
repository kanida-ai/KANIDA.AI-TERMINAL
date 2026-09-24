"""Falcon Conviction-Tier classifier (Enterprise / Gold / Tier 2 / Avoid).

The doc-validated 3-gate tier layer (G1 entry-day confirmation, G2 no 3-sector
cluster, G3 not-a-serial-loser) applied to the LIVE Falcon Top 10 output.

This is the REVIEW build: read-only on PROD + RND, writes an Excel for operator
sign-off. After review the logic ports into a power_user service for production.

Naming: "conviction tier" (Enterprise/Gold/Tier2/Avoid) — deliberately distinct
from the existing live_tier.py rank-buckets (ELITE/HIGH/MID/...), which are a
SEPARATE intraday-entry system. Do not conflate.

Data sources (all READ-ONLY):
  PROD  data/db/kanida_universe.db
    falcon_signals_live   — daily picks (symbol, sector, n_fires, score, entry_date)
    ohlc_daily            — entry-day open/close for D1 (confirmed tier, past days)
    falcon_sectors        — sector backfill
  RND   universe_engine/data/db/kanida_universe.db
    falcon_signal_day_study (persona falcon_top10_daily) — resolved top-10 picks
                            for the rolling-90d serial-loser computation

Two-stage by construction:
  * Latest signal_date (entry = next trading day, not yet elapsed) → PROVISIONAL
    candidates only (G1/D1 unknown until entry-day 3:30 PM).
  * Previous signal_date (entry elapsed, ohlc present) → CONFIRMED tier (D1 known).

Usage:
  python falcon_tier_classifier.py --prod <PROD.db> --rnd <RND.db> --out <xlsx>
"""
from __future__ import annotations

import argparse
import sqlite3
import sys
from datetime import date, datetime, timedelta
from typing import Any, Dict, List, Optional, Tuple

sys.stdout.reconfigure(encoding="utf-8")

MIN_FIRES = 10          # Falcon Top 10 locked persona gate (PERSONA_CONFIGS)
TOP_N = 10
SERIAL_WINDOW_DAYS = 90
SERIAL_MIN_APPEAR = 5
SERIAL_AVG_MAX = 0.0    # avg net return < 0
SERIAL_WR_MAX = 0.35    # win rate < 35%
ENTERPRISE_D1 = 10.0
G1_D1 = 0.5
D1_EXIT = -2.0
RND_PERSONA = "falcon_top10_daily"

# recent_trend_warning: a SEPARATE, NON-BLOCKING caution flag (operator decision
# 2026-06-16). G3 (90-day serial loser) stays strict + unchanged. This catches a
# stock that is CURRENTLY bleeding even though its 90-day average is positive
# (e.g. HFCL: +5.8% / 44% WR over 90d, rescued by an April pump, but on a current
# losing streak). Intent was "last 6 all negative", but the data showed HFCL's
# last 6 are NOT all negative (a +20.5% bounce on 2026-05-25 sits in that window).
# A current consecutive-losing-STREAK captures the intent correctly. >=3 fires on
# HFCL today; >=4 does not. Non-blocking: downgrades Gold->MANUAL_REVIEW, never AVOID.
RECENT_STREAK_N = 3


# ── data access ─────────────────────────────────────────────────────────
def latest_signal_dates(prod: sqlite3.Connection, n: int = 2) -> List[str]:
    rows = prod.execute(
        "SELECT DISTINCT signal_date FROM falcon_signals_live "
        "WHERE signal_date IS NOT NULL ORDER BY signal_date DESC LIMIT ?", (n,)
    ).fetchall()
    return [r[0] for r in rows]


def top10(prod: sqlite3.Connection, signal_date: str) -> List[Dict[str, Any]]:
    """Top-10 for the day — EXACTLY the portal ranker: avg_lift = score/n_fires,
    n_fires >= MIN_FIRES, all500 universe, DESC, LIMIT 10."""
    rows = prod.execute(
        """
        SELECT symbol, sector, n_fires, score, entry_date, close_at_signal,
               (score * 1.0 / n_fires) AS avg_lift
          FROM falcon_signals_live
         WHERE signal_date = ? AND n_fires >= ?
         ORDER BY avg_lift DESC
         LIMIT ?
        """, (signal_date, MIN_FIRES, TOP_N)
    ).fetchall()
    sectors = dict(prod.execute("SELECT symbol, sector FROM falcon_sectors").fetchall())
    out = []
    for i, (sym, sec, nf, sc, ed, cas, al) in enumerate(rows, start=1):
        out.append({
            "rank": i, "symbol": sym, "sector": sec or sectors.get(sym) or "",
            "n_fires": nf, "score": sc, "entry_date": ed, "close_at_signal": cas,
            "avg_lift": al,
        })
    return out


def entry_day_d1(prod: sqlite3.Connection, symbol: str, entry_date: str
                 ) -> Optional[float]:
    """D1 close return = (entry-day close - entry-day open)/open * 100.
    Entry is at the entry_date 9:15 open; D1 is that same day's 3:30 close.
    None if ohlc not yet present (entry day not elapsed)."""
    row = prod.execute(
        "SELECT open, close FROM ohlc_daily WHERE symbol=? AND trade_date=?",
        (symbol, entry_date)
    ).fetchone()
    if not row or row[0] in (None, 0):
        return None
    o, c = row
    return (c / o - 1.0) * 100.0


def serial_losers(rnd: sqlite3.Connection, as_of: str
                  ) -> Tuple[set, Dict[str, Dict[str, Any]]]:
    """Serial-loser set as-of `as_of` (forward-only): symbols whose RESOLVED
    top-10 appearances in the prior 90 calendar days satisfy ALL THREE:
    n>=5, avg net_ret<0, win_rate<35%. Returns (set, per-symbol stats for ALL
    symbols with >=5 appearances so the operator can see near-misses)."""
    d = datetime.strptime(as_of, "%Y-%m-%d").date()
    lo = (d - timedelta(days=SERIAL_WINDOW_DAYS)).isoformat()
    rows = rnd.execute(
        """
        SELECT symbol, net_ret_pct
          FROM falcon_signal_day_study
         WHERE persona = ?
           AND signal_date >= ? AND signal_date < ?
           AND exit_date IS NOT NULL AND exit_date < ?
           AND net_ret_pct IS NOT NULL
        """, (RND_PERSONA, lo, as_of, as_of)
    ).fetchall()
    by: Dict[str, List[float]] = {}
    for sym, nr in rows:
        by.setdefault(sym, []).append(float(nr))
    losers: set = set()
    stats: Dict[str, Dict[str, Any]] = {}
    for sym, rets in by.items():
        n = len(rets)
        if n < SERIAL_MIN_APPEAR:
            continue
        avg = sum(rets) / n
        wr = sum(1 for x in rets if x > 0) / n
        is_loser = (avg < SERIAL_AVG_MAX and wr < SERIAL_WR_MAX)
        stats[sym] = {"n": n, "avg": avg, "wr": wr * 100, "is_loser": is_loser}
        if is_loser:
            losers.add(sym)
    return losers, stats


def recent_losing_streak(rnd: sqlite3.Connection, symbol: str, as_of: str) -> int:
    """Current consecutive losing streak: count of the most-recent RESOLVED
    appearances (signal_date < as_of, exit_date < as_of) with net_ret_pct < 0,
    stopping at the first win. Drives the non-blocking recent_trend_warning."""
    rows = rnd.execute(
        """
        SELECT net_ret_pct FROM falcon_signal_day_study
         WHERE persona = ? AND symbol = ? AND signal_date < ?
           AND exit_date IS NOT NULL AND exit_date < ? AND net_ret_pct IS NOT NULL
         ORDER BY signal_date DESC, exit_date DESC
        """, (RND_PERSONA, symbol, as_of, as_of)
    ).fetchall()
    streak = 0
    for (nr,) in rows:
        if nr < 0:
            streak += 1
        else:
            break
    return streak


# ── classification ──────────────────────────────────────────────────────
def classify_day(prod, rnd, signal_date: str) -> Tuple[List[Dict[str, Any]], Dict, set, Dict]:
    picks = top10(prod, signal_date)
    if not picks:
        return [], {}, set(), {}
    # G2: sector cluster count within this day's top-10
    sec_count: Dict[str, int] = {}
    for p in picks:
        sec_count[p["sector"]] = sec_count.get(p["sector"], 0) + 1
    losers, lstats = serial_losers(rnd, signal_date)
    entry_date = picks[0]["entry_date"]
    # is this day's entry elapsed? (ohlc present for any pick on entry_date)
    confirmed = any(entry_day_d1(prod, p["symbol"], entry_date) is not None for p in picks)

    out = []
    for p in picks:
        sym = p["symbol"]
        sc = sec_count.get(p["sector"], 1)
        g2 = sc <= 2
        g3 = sym not in losers
        sl = lstats.get(sym)
        sl_reason = ""
        if not g3 and sl:
            sl_reason = f"serial loser: n={sl['n']}, avg={sl['avg']:+.1f}%, WR={sl['wr']:.0f}%"
        d1 = entry_day_d1(prod, sym, entry_date) if confirmed else None
        g1 = (d1 is not None and d1 > G1_D1)
        cluster_warn = not g2
        streak = recent_losing_streak(rnd, sym, signal_date)
        recent_trend_warning = streak >= RECENT_STREAK_N

        # ── PROVISIONAL (always computable at signal time: G2 + G3 + rank) ──
        if not g3:
            prov = "AVOID"
        elif p["rank"] <= 3 and g2:
            prov = "ENTERPRISE_CANDIDATE"
        else:
            prov = "GOLD_CANDIDATE" + (" (cluster)" if cluster_warn else "")

        # ── CONFIRMED (only when entry day elapsed → D1 known) ──
        conf = "PENDING (D1 at entry-day 3:30 PM)"
        action = ""
        if confirmed and d1 is not None:
            if not g3:
                conf = "AVOID"
                action = "Block — serial loser"
            elif p["rank"] <= 3 and d1 > ENTERPRISE_D1 and g3:
                conf = "ENTERPRISE"
                action = "Full size ₹50k — highest conviction"
            elif g1 and g3:  # Gold (cluster-fail still Gold + warning per doc)
                conf = "GOLD" + (" (cluster_warning)" if cluster_warn else "")
                action = "Full size ₹50k"
            elif g3 and d1 < D1_EXIT:
                conf = "TIER2 / D1-EXIT"
                action = f"Exit-evening recommended (D1 {d1:+.1f}% < -2%)"
            else:  # G3 true, D1 in (-2%, +0.5%]
                conf = "TIER2"
                action = "Manual review, half size ₹25k"
        elif not confirmed:
            if not g3:
                action = "Block tomorrow — serial loser"
            elif prov == "ENTERPRISE_CANDIDATE":
                action = "Watch for D1>+10% tomorrow → Enterprise"
            else:
                action = "Watch for D1>+0.5% tomorrow → Gold"

        # ── recent_trend_warning: NON-BLOCKING. G3 untouched. Downgrades a
        #    Gold/Gold-candidate to MANUAL_REVIEW (amber), never to AVOID, and
        #    leaves ENTERPRISE alone (a fresh +10% entry day overrides recent
        #    weakness). AVOID stays AVOID. ──
        eff_tier_field = "confirmed_tier" if (confirmed and d1 is not None) else "provisional_tier"
        cur_tier = conf if (confirmed and d1 is not None) else prov
        if recent_trend_warning and "AVOID" not in cur_tier and "ENTERPRISE" not in cur_tier:
            if "GOLD" in cur_tier:
                cur_tier = "MANUAL_REVIEW (recent_trend_warning)"
            action = (f"⚠ recent losing streak {streak} — MANUAL_REVIEW; "
                      + (action or "")).strip()
            if eff_tier_field == "confirmed_tier":
                conf = cur_tier
            else:
                prov = cur_tier
        elif recent_trend_warning:
            action = (f"⚠ recent losing streak {streak}. " + (action or "")).strip()

        out.append({
            "signal_date": signal_date, "entry_date": entry_date,
            "rank": p["rank"], "symbol": sym, "sector": p["sector"],
            "avg_lift": round(p["avg_lift"], 2), "n_fires": p["n_fires"],
            "sector_count_today": sc,
            "G2_no_cluster": g2, "G3_not_serial_loser": g3,
            "serial_loser_reason": sl_reason,
            "d1_close_ret": (None if d1 is None else round(d1, 2)),
            "G1_d1_green": (None if d1 is None else g1),
            "provisional_tier": prov,
            "confirmed_tier": conf,
            "cluster_warning": cluster_warn,
            "recent_losing_streak": streak,
            "recent_trend_warning": recent_trend_warning,
            "action": action,
        })
    meta = {"signal_date": signal_date, "entry_date": entry_date,
            "confirmed": confirmed, "n_losers": len(losers)}
    return out, meta, losers, lstats


# ── excel ───────────────────────────────────────────────────────────────
def write_excel(path: str, days: List[Tuple[List[Dict], Dict, set, Dict]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "Tier Classification"
    hdr_fill = PatternFill("solid", fgColor="1F2937")
    hdr_font = Font(color="FFFFFF", bold=True)
    tier_fill = {
        "ENTERPRISE": "C6EFCE", "ENTERPRISE_CANDIDATE": "E2EFDA",
        "GOLD": "FFF2CC", "GOLD_CANDIDATE": "FFF9E6",
        "TIER2": "FCE4D6", "AVOID": "FFC7CE",
        "MANUAL_REVIEW": "FFE699",   # amber — recent_trend_warning downgrade
    }
    cols = ["signal_date", "entry_date", "rank", "symbol", "sector", "avg_lift",
            "n_fires", "sector_count_today", "G2_no_cluster", "G3_not_serial_loser",
            "serial_loser_reason", "d1_close_ret", "G1_d1_green",
            "provisional_tier", "confirmed_tier", "cluster_warning",
            "recent_losing_streak", "recent_trend_warning", "action"]

    r = 1
    for picks, meta, losers, lstats in days:
        label = ("LATEST — PROVISIONAL (entry %s not yet elapsed)" % meta["entry_date"]
                 if not meta["confirmed"]
                 else "PREVIOUS — CONFIRMED (entry %s elapsed, D1 known)" % meta["entry_date"])
        ws.cell(r, 1, f"Signal {meta['signal_date']}  |  {label}  |  serial losers active: {meta['n_losers']}")
        ws.cell(r, 1).font = Font(bold=True, size=12)
        r += 1
        for ci, c in enumerate(cols, 1):
            cell = ws.cell(r, ci, c)
            cell.fill = hdr_fill; cell.font = hdr_font
        r += 1
        for p in picks:
            for ci, c in enumerate(cols, 1):
                v = p.get(c)
                ws.cell(r, ci, "" if v is None else v)
            key = (p["confirmed_tier"].split()[0] if p["confirmed_tier"] and "PENDING" not in p["confirmed_tier"]
                   else p["provisional_tier"].split()[0])
            fill = tier_fill.get(key)
            if fill:
                for ci in range(1, len(cols) + 1):
                    ws.cell(r, ci).fill = PatternFill("solid", fgColor=fill)
            r += 1
        r += 2

    # widths
    widths = [12, 11, 5, 12, 22, 9, 8, 10, 8, 10, 34, 11, 8, 22, 34, 9, 9, 10, 44]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w

    # ── Serial-loser transparency sheet ──
    ws2 = wb.create_sheet("Serial-Loser Window (90d)")
    ws2.append(["As-of signal_date", "symbol", "appearances_90d", "avg_net_ret_%",
                "win_rate_%", "flagged_serial_loser", "note"])
    for ci in range(1, 8):
        ws2.cell(1, ci).fill = hdr_fill; ws2.cell(1, ci).font = hdr_font
    for picks, meta, losers, lstats in days:
        for sym, s in sorted(lstats.items(), key=lambda kv: kv[1]["avg"]):
            note = ""
            if not s["is_loser"]:
                fails = []
                if s["n"] < SERIAL_MIN_APPEAR: fails.append("n<5")
                if not s["avg"] < 0: fails.append("avg not<0")
                if not s["wr"] < 35: fails.append("WR not<35%")
                note = "not flagged: " + ", ".join(fails)
            ws2.append([meta["signal_date"], sym, s["n"], round(s["avg"], 2),
                        round(s["wr"], 1), bool(s["is_loser"]), note])
    for i, w in enumerate([16, 12, 16, 14, 12, 20, 36], 1):
        ws2.column_dimensions[ws2.cell(1, i).column_letter].width = w

    # ── Notes sheet ──
    ws3 = wb.create_sheet("Notes & Rules")
    notes = [
        "FALCON CONVICTION-TIER CLASSIFICATION — REVIEW BUILD",
        "",
        "Gates:",
        "  G1 = entry-day close return > +0.5% (the D1 move). Known only AFTER entry day's 3:30 PM.",
        "  G2 = <=2 stocks from the same sector in that day's top-10 (no 3-cluster).",
        "  G3 = NOT a serial loser. Serial loser = over the prior 90 days of RESOLVED top-10 picks,",
        "       ALL THREE: appearances>=5 AND avg net return<0 AND win rate<35%.",
        "",
        "Tiers (confirmed, once D1 is known):",
        "  ENTERPRISE = rank<=3 AND D1>+10% AND G3.   GOLD = G1 AND G3 (cluster-fail => Gold+warning).",
        "  TIER2 = G3 AND D1<=+0.5% (D1<-2% => exit-evening).   AVOID = serial loser (G3 False).",
        "",
        "TWO STAGES (why the two sheets differ):",
        "  LATEST signal day: entry is the NEXT trading day, not yet elapsed -> D1/G1 unknown ->",
        "    only PROVISIONAL candidates (Enterprise-candidate / Gold-candidate / Avoid) can be set.",
        "  PREVIOUS signal day: entry has elapsed, ohlc present -> CONFIRMED tier with real D1.",
        "",
        "recent_trend_warning (NON-BLOCKING, operator decision 2026-06-16):",
        "  G3 (90-day serial loser) stays STRICT + UNCHANGED — it governs blocking (AVOID).",
        "  recent_trend_warning is a SEPARATE caution flag: current consecutive losing streak >= 3",
        "  (most-recent resolved appearances all negative, stopping at the first win). It does NOT",
        "  block; it downgrades a Gold / Gold-candidate to MANUAL_REVIEW and shows an amber note.",
        "  ENTERPRISE is left alone (a fresh +10% entry day overrides recent weakness); AVOID stays AVOID.",
        "  NOTE: the original intent was 'last 6 appearances all negative', but HFCL's last 6 are NOT",
        "  all negative — a +20.5% bounce on 2026-05-25 sits in that window (90d avg +5.8%, WR 44%).",
        "  A current losing STREAK>=3 captures the intent and fires on HFCL today (streak=3); >=4 would not.",
        "  Threshold RECENT_STREAK_N is a single constant — change it in one place if you prefer 4.",
        "",
        "Return basis: D1 and tier returns are measured from the entry-day 9:15 OPEN (entry price).",
        "Sources: PROD falcon_signals_live + ohlc_daily (read-only); RND falcon_signal_day_study (serial losers).",
        "This is a REVIEW artifact. Production wiring happens only after sign-off.",
    ]
    for i, line in enumerate(notes, 1):
        ws3.cell(i, 1, line)
        if i == 1:
            ws3.cell(i, 1).font = Font(bold=True, size=13)
    ws3.column_dimensions["A"].width = 110

    wb.save(path)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--prod", required=True)
    ap.add_argument("--rnd", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    prod = sqlite3.connect(a.prod); prod.row_factory = None
    rnd = sqlite3.connect(a.rnd)
    sigs = latest_signal_dates(prod, 2)
    print(f"[tier] signal dates: {sigs}")
    days = []
    for sd in sigs:                          # [latest, previous]
        picks, meta, losers, lstats = classify_day(prod, rnd, sd)
        days.append((picks, meta, losers, lstats))
        stage = "CONFIRMED" if meta.get("confirmed") else "PROVISIONAL"
        print(f"\n[tier] {sd}  ({stage}, entry {meta.get('entry_date')}, "
              f"serial-losers={meta.get('n_losers')})")
        for p in picks:
            tier = p["confirmed_tier"] if meta["confirmed"] else p["provisional_tier"]
            d1 = "" if p["d1_close_ret"] is None else f"D1={p['d1_close_ret']:+.2f}%"
            print(f"   #{p['rank']:2} {p['symbol']:12} {str(p['sector'])[:20]:20} "
                  f"avgL={p['avg_lift']:5.2f} {d1:12} G2={int(p['G2_no_cluster'])} "
                  f"G3={int(p['G3_not_serial_loser'])} -> {tier}")
    write_excel(a.out, days)
    print(f"\n[tier] wrote {a.out}")
    prod.close(); rnd.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
