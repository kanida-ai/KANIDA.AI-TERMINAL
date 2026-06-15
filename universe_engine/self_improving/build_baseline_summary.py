#!/usr/bin/env python3
"""build_baseline_summary.py — generate falcon_baseline_summary.xlsx from the
populated falcon_baseline_trades table (Steps 2+3). READ-ONLY on the DB.

IMPORTANT — two different "returns":
  * Walk-forward return (the LOCKED parity figure, e.g. 2022 +76.36%) is
    EQUITY-BASED: end_equity/₹5L, INCLUDING the year-end mark-to-market of
    positions still open at year end. That number is produced by the engine
    (build_baseline.py) and verified in Step 2 — it CANNOT be reconstructed from
    this table because open-at-end trades are stored unresolved (net_pnl=0,
    exit_reason=NULL), which is correct for the learning layer.
  * Closed-trade return (shown here, derived) = sum(net_pnl of CLOSED trades)/₹5L.
    This EXCLUDES open-position MTM, so it is lower than the walk-forward figure.

This workbook therefore shows the faithful closed-trade aggregates + full trade
detail/journey/sector-attribution from the table, and lists the verified
walk-forward returns (Step 2) as a labelled reference — it does not re-derive them.
"""
from __future__ import annotations
import argparse, sqlite3, sys
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

CASH_START = 500_000.0
PERSONA = "falcon_top10"
CLOSED = ("TIME_STOP", "INIT_STOP", "TRAIL_GIVEBACK")
# Verified Step-2 walk-forward returns (equity-based, incl. open-MTM). 2021-2025
# are the LOCKED parity numbers; 2026 is the partial-year value from the same run.
WALKFWD = {2021: 5.35, 2022: 76.36, 2023: 564.47, 2024: 469.05, 2025: 346.64, 2026: 86.03}
LOCKED_N = {2021: 15, 2022: 317, 2023: 810, 2024: 846, 2025: 644}


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--rnd-db", required=True)
    ap.add_argument("--out", default=str(Path(__file__).parent / "out" / "falcon_baseline_summary.xlsx"))
    a = ap.parse_args()

    con = sqlite3.connect(f"file:{a.rnd_db}?mode=ro", uri=True); con.row_factory = sqlite3.Row
    rows = con.execute("SELECT * FROM falcon_baseline_trades WHERE persona=? ORDER BY signal_date, symbol", (PERSONA,)).fetchall()
    con.close()
    if not rows:
        print("no baseline rows — run build_baseline.py first"); return 1

    agg = {}
    for r in rows:
        y = int(str(r["signal_date"])[:4])
        d = agg.setdefault(y, dict(n=0, closed=0, openend=0, wins=0, retsum=0.0, cpnl=0.0, bw=0, bl=0))
        d["n"] += 1
        if r["exit_reason"] in CLOSED:
            d["closed"] += 1; d["cpnl"] += r["net_pnl"] or 0.0
            if (r["net_ret_pct"] or 0) > 0: d["wins"] += 1
            d["retsum"] += r["net_ret_pct"] or 0.0
        else:
            d["openend"] += 1
        if r["big_winner_flag"] == 1: d["bw"] += 1
        if r["big_loser_flag"] == 1: d["bl"] += 1

    print(f"\n{'Yr':>4} {'nClosed':>7} {'nOpen':>5} {'closedRet%':>10} {'walkFwd%(verified)':>18} {'win%':>6} {'bigW':>5} {'bigL':>5}")
    for y in sorted(agg):
        d = agg[y]; cret = d["cpnl"] / CASH_START * 100
        wr = d["wins"] / d["closed"] * 100 if d["closed"] else 0
        print(f"{y:>4} {d['closed']:>7} {d['openend']:>5} {cret:>10.2f} {str(WALKFWD.get(y,'—')):>18} {wr:>6.1f} {d['bw']:>5} {d['bl']:>5}")

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        print("openpyxl missing"); return 2
    wb = Workbook(); ws = wb.active; ws.title = "Per-Year Summary"
    hdr = ["Year", "Walk-forward Return% (verified Step-2, equity-based incl. open-MTM)",
           "Locked n_closed", "n_closed (table)", "n_open_at_end",
           "Closed-trade Return% (derived, excl. open-MTM)", "win_rate%",
           "avg_closed_ret%", "big_winners", "big_losers", "closed_only_pnl"]
    ws.append(hdr); [setattr(c, "font", Font(bold=True)) for c in ws[1]]
    for y in sorted(agg):
        d = agg[y]; cret = round(d["cpnl"] / CASH_START * 100, 2)
        wr = round(d["wins"] / d["closed"] * 100, 2) if d["closed"] else 0
        avg = round(d["retsum"] / d["closed"], 2) if d["closed"] else 0
        ws.append([y, WALKFWD.get(y), LOCKED_N.get(y), d["closed"], d["openend"],
                   cret, wr, avg, d["bw"], d["bl"], round(d["cpnl"])])
    ws.append([])
    ws.append(["NOTE: 'Walk-forward Return%' is the equity-based parity figure verified in Step 2 "
               "(end_equity/₹5L, includes year-end MTM of open positions). The 'Closed-trade Return%' "
               "here is derived from this table and is LOWER because open-at-end positions are stored "
               "unresolved (net_pnl=0) — correct for the learning layer, which only uses resolved trades. "
               "2021-2025 walk-forward = locked; 2026 = partial year."])

    ws2 = wb.create_sheet("All Trades")
    keep = ["signal_date","entry_date","exit_date","symbol","sector","engine_rank","avg_lift","n_fires",
            "entry_price","exit_price","exit_reason","hold_days_trading","shares","net_pnl","net_ret_pct",
            "peak_ret_during_hold","peak_day_during_hold","trough_ret_during_hold","peak_sustained",
            "big_winner_flag","big_loser_flag","move_type","sector_ret_same_period","stock_vs_sector"]
    ws2.append(keep); [setattr(c, "font", Font(bold=True)) for c in ws2[1]]
    for r in rows:
        ws2.append([r[k] if k in r.keys() else None for k in keep])

    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    wb.save(a.out)
    print(f"\nwrote {a.out}  ({len(rows)} trades)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
