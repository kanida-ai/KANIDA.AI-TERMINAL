"""Validated full report — RECOMMENDED config = optimized exits + REGIME GATE
(trade only when the strategy's trailing-10-day avg daily return is positive;
warm-up first 10 days = trade, no basis to skip). Same validation + traceable
logs as the ungated report, plus transparent Traded/Skipped columns. New file.
"""
import sys, pickle
sys.path.insert(0, r"C:\Users\SPS\Desktop\Kanida.ai Terminal Quant Intelligence Engine\scripts")
from pathlib import Path
from collections import defaultdict
import numpy as np
from falcon_full_report import (ROOT, ASSIGNED, CHARGE_SIDE, MONTHS, RND_1MIN_MAX,
                                validate, basket_exit_q, month_status)
from falcon_stock_journal_opt import ARM, FLOOR, GIVE, STOP

GATE_WINDOW = 10


def main():
    mats = pickle.load(open(ROOT / "docs" / "ops" / "_opt_dataset.pkl", "rb"))
    dates = [d for d, _ in mats]
    gt_ext, gaps, muhurat = validate(dates)
    print("=" * 70)
    print("DATA VALIDATION (ground truth = actual 1-min market data)")
    print(f"  dataset days: {len(dates)} ({dates[0]} -> {dates[-1]})")
    print(f"  Muhurat/special excluded: {muhurat}")
    print("  REAL GAPS: none" if not gaps else f"  ** REAL GAPS: {gaps} **")

    gt_norm = [d for d in gt_ext if d not in set(muhurat)]
    exp_by_month = defaultdict(int)
    for d in gt_norm:
        exp_by_month[d[:7]] += 1
    mn, mx = dates[0], dates[-1]

    # ---- PASS 1: compute the full trade for EVERY day (rupee-consistent at Rs5L) ----
    dc = []
    for d, M in mats:
        n = M["n"]; entry = M["entry"]; close = M["close"]; opn = M["opn"]; high = M["high"]; grid = M["grid"]
        alloc = ASSIGNED / M["nstocks"]
        qty = np.floor(alloc / entry)
        deployed = float((entry * qty).sum())
        xb, reason = basket_exit_q(close, qty, deployed, n)
        if reason == "EOD":
            exit_px = close[n - 1].astype(float); exit_time = grid[n - 1]
        else:
            nx = opn[xb + 1].astype(float)
            exit_px = np.where(np.isfinite(nx) & (nx > 0), nx, close[xb]).astype(float)
            exit_time = grid[min(xb + 1, n - 1)]
        exit_val = float((exit_px * qty).sum())
        gross = exit_val - deployed
        charges = (deployed + exit_val) * CHARGE_SIDE
        net = gross - charges
        mfe = float((((high - entry) * qty).sum(axis=1) / deployed * 100).max())
        stocks = []
        for j in range(M["nstocks"]):
            e = float(entry[j]); q = int(qty[j]); xp = float(exit_px[j])
            stocks.append(dict(rank=M["rank"][j], symbol=M["syms"][j], entry_price=round(e, 2),
                               alloc=round(alloc, 0), qty=q, deployed=round(q * e, 0),
                               exit_price=round(xp, 2), gross_pnl=round(q * (xp - e), 0),
                               stock_ret=round((xp / e - 1) * 100, 3)))
        dc.append(dict(date=d, month=d[:7], n_traded=M["nstocks"], deployed=round(deployed, 0),
                       undeployed=round(ASSIGNED - deployed, 0), exit_time=exit_time, exit_reason=reason,
                       gross=round(gross, 0), charges=round(charges, 0), net=round(net, 0),
                       gret=round(gross / deployed * 100, 3), nret=round(net / deployed * 100, 3),
                       gret_asg=round(gross / ASSIGNED * 100, 3), mfe=mfe, stocks=stocks))

    # ---- REGIME GATE (walk-forward; warm-up trades) ----
    grets = [x["gret"] for x in dc]
    for i, x in enumerate(dc):
        x["traded"] = True if i < GATE_WINDOW else (np.mean(grets[i - GATE_WINDOW:i]) > 0)
        x["gate_trail10"] = round(float(np.mean(grets[i - GATE_WINDOW:i])), 3) if i >= GATE_WINDOW else None

    n_traded = sum(1 for x in dc if x["traded"])
    print(f"  REGIME GATE: traded {n_traded} / {len(dc)} days ({len(dc)-n_traded} skipped)")

    write_xlsx(dc, muhurat, gaps, mn, mx, exp_by_month)
    print(f"[*] validation PASS={'YES' if not gaps else 'NO'}")


def write_xlsx(dc, muhurat, gaps, mn, mx, exp_by_month):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    F = "Calibri"
    TITLE = Font(name=F, bold=True, size=13, color="1F6F8B"); NOTE = Font(name=F, size=9, italic=True, color="666666")
    H = Font(name=F, bold=True, size=9, color="FFFFFF"); HDR = PatternFill("solid", start_color="1F6F8B")
    NORM = Font(name=F, size=10); GRN = Font(name=F, size=10, color="1E7E34"); RED = Font(name=F, size=10, color="B00020")
    GRNF = PatternFill("solid", start_color="E7F4EA"); REDF = PatternFill("solid", start_color="FDECEA")
    AMBF = PatternFill("solid", start_color="FFF4CE"); GRYF = PatternFill("solid", start_color="EEEEEE")
    thin = Side(style="thin", color="DDDDDD"); BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
    Cc = Alignment(horizontal="center"); Rr = Alignment(horizontal="right"); Ll = Alignment(horizontal="left")
    RS = '#,##0;[Red]-#,##0'
    wb = Workbook()

    # monthly aggregation (traded-day metrics; fresh Rs5L/month simple, net)
    by = defaultdict(list)
    for x in dc:
        by[x["month"]].append(x)
    monthly = []
    for ym in sorted(by):
        rows = by[ym]
        traded = [r for r in rows if r["traded"]]
        nret = np.array([r["nret"] for r in traded]) if traded else np.array([])
        gret = np.array([r["gret"] for r in traded]) if traded else np.array([])
        nstk = np.array([r["n_traded"] for r in traded]) if traded else np.array([])
        wins = nret[nret > 0] if len(nret) else np.array([]); losses = nret[nret < 0] if len(nret) else np.array([])
        sum_net = float(nret.sum()) if len(nret) else 0.0
        pnl = ASSIGNED * sum_net / 100.0
        monthly.append(dict(
            year=int(ym[:4]), month=int(ym[5:7]), days=len(rows), traded=len(traded), skip=len(rows) - len(traded),
            avgstk=round(float(nstk.mean()), 2) if len(nstk) else 0.0,
            win=int((nret > 0).sum()), loss=int((nret < 0).sum()),
            winp=round(float((nret > 0).mean() * 100), 1) if len(nret) else 0.0,
            avgwin=round(float(wins.mean()), 2) if len(wins) else 0.0,
            avgloss=round(float(losses.mean()), 2) if len(losses) else 0.0,
            sum_gross=round(float(gret.sum()), 1) if len(gret) else 0.0,
            d1=int((nret >= 1).sum()) if len(nret) else 0,
            t05=sum(1 for r in traded if r["mfe"] >= 0.5),
            mret=round(sum_net, 2), pnl=round(pnl, 0), mend=round(ASSIGNED + pnl, 0),
            status=month_status(ym, len(rows), exp_by_month, mn, mx, muhurat)))

    # ---- Validation ----
    ws = wb.active; ws.title = "Validation"
    ws.cell(1, 1, "Data Validation — RECOMMENDED config (optimized exits + regime gate)").font = TITLE
    ws.cell(2, 1, "Ground truth = actual 1-min market data. 'Days' = real trading days; 'Traded'/'Skipped' = after the regime gate. "
                  "Muhurat/special sessions (no 09:15) excluded. GATE: trade only when the strategy's trailing-10-day avg is positive (first 10 days trade).").font = NOTE
    ws.cell(4, 1, "OVERALL: " + ("PASS — no real gaps" if not gaps else f"FAIL — gaps {gaps}")).font = Font(name=F, bold=True, size=11, color=("1E7E34" if not gaps else "B00020"))
    ws.cell(5, 1, f"Muhurat/special excluded: {muhurat or 'none'}").font = NORM
    for c, h in enumerate(["Month", "Trading days", "Traded (gate on)", "Skipped (gate off)", "Status"], 1):
        cell = ws.cell(7, c, h); cell.font = H; cell.fill = HDR; cell.border = BORD; cell.alignment = Cc
    r = 8
    for m in monthly:
        ym = f"{m['year']}-{m['month']:02d}"
        vals = [ym, m["days"], m["traded"], m["skip"], m["status"]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v); cell.font = NORM; cell.border = BORD; cell.alignment = Ll if c in (1, 5) else Rr
        fill = GRNF if m["status"] == "COMPLETE" else (AMBF if "PARTIAL" in m["status"] else REDF)
        for c in range(1, 6): ws.cell(r, c).fill = fill
        r += 1
    for i, w in enumerate([10, 14, 18, 18, 32], 1): ws.column_dimensions[get_column_letter(i)].width = w

    # ---- Monthly Summary ----
    ws2 = wb.create_sheet("Monthly Summary")
    ws2.cell(1, 1, "Monthly Summary — RECOMMENDED config (optimized exits + regime gate), CASH Rs5L").font = TITLE
    ws2.cell(2, 1, "Performance columns are on TRADED days (regime on). Sum_month% = GROSS; Month Ret% = NET (after ~0.10%/day). "
                   "Returns on DEPLOYED capital. Month P&L = Rs5L x Month Ret%(net), fresh each month (no compounding).").font = NOTE
    heads2 = ["Year", "Month", "Days", "Traded", "Skipped", "Avg_stk", "Win", "Loss", "Win%", "Avg_win%", "Avg_loss%",
              "Sum_month%(gross)", "Days>=1%", ">+0.5% touch", "Month Ret%(net)", "Month P&L Rs", "Month-End Rs", "Grew/Eroded", "Status"]
    for c, h in enumerate(heads2, 1):
        cell = ws2.cell(4, c, h); cell.font = H; cell.fill = HDR; cell.border = BORD; cell.alignment = Cc
    r = 5
    for m in monthly:
        arrow = "GREW" if m["pnl"] >= 0 else "ERODED"
        vals = [m["year"], MONTHS[m["month"]], m["days"], m["traded"], m["skip"], m["avgstk"], m["win"], m["loss"],
                m["winp"], m["avgwin"], m["avgloss"], m["sum_gross"], m["d1"], m["t05"], m["mret"], m["pnl"], m["mend"], arrow, m["status"]]
        for c, v in enumerate(vals, 1):
            cell = ws2.cell(r, c, v); cell.font = NORM; cell.border = BORD; cell.alignment = Cc if c in (1, 2, 18) else Rr
            if c in (16, 17): cell.number_format = RS
            if c == 9: cell.number_format = '0.0"%"'
            if c in (10, 11, 12, 15): cell.number_format = '0.00'
            if c == 16: cell.font = GRN if v >= 0 else RED
        r += 1
    for i, w in enumerate([6, 6, 6, 7, 8, 8, 5, 5, 7, 9, 10, 16, 9, 12, 14, 13, 14, 11, 24], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "C5"

    # ---- Daily Trade Log (ALL days; skipped rows marked) ----
    ws3 = wb.create_sheet("Daily Trade Log")
    h3 = ["date", "regime", "trail10_signal", "stocks_traded", "assigned_cap", "deployed", "undeployed_cash",
          "exit_time", "exit_reason", "gross_pnl", "charges", "net_pnl", "return_denominator",
          "gross_ret_%_dep", "NET_ret_%_dep"]
    for c, h in enumerate(h3, 1):
        cell = ws3.cell(1, c, h); cell.font = H; cell.fill = HDR; cell.border = BORD; cell.alignment = Cc
    ri = 2
    for x in dc:
        if x["traded"]:
            vals = [x["date"], "TRADE", x["gate_trail10"], x["n_traded"], ASSIGNED, x["deployed"], x["undeployed"],
                    x["exit_time"], x["exit_reason"], x["gross"], x["charges"], x["net"], "deployed capital",
                    x["gret"], x["nret"]]
        else:
            vals = [x["date"], "SKIP-REGIME", x["gate_trail10"], 0, ASSIGNED, 0, ASSIGNED, "-", "not traded (regime off)",
                    0, 0, 0, "n/a (cash)", 0.0, 0.0]
        for c, v in enumerate(vals, 1):
            cell = ws3.cell(ri, c, v); cell.font = NORM; cell.alignment = Ll if c in (1, 2, 8, 9, 13) else Rr
            if c in (5, 6, 7, 10, 11, 12): cell.number_format = RS
            if c in (3, 14, 15): cell.number_format = '0.000'
        fill = GRYF if not x["traded"] else (GRNF if x["net"] > 0 else (REDF if x["net"] < 0 else None))
        if fill:
            for c in range(1, 16): ws3.cell(ri, c).fill = fill
        ri += 1
    for i, w in enumerate([11, 12, 14, 13, 12, 11, 15, 9, 22, 11, 10, 11, 18, 14, 14], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    ws3.freeze_panes = "A2"; ws3.auto_filter.ref = f"A1:O{len(dc)+1}"

    # ---- Per-Stock Log (traded days only) ----
    ws4 = wb.create_sheet("Per-Stock Log")
    h4 = ["date", "rank", "symbol", "entry_time", "entry_price", "alloc", "qty=floor(alloc/entry)",
          "deployed=qty*entry", "exit_time", "exit_price", "exit_reason", "gross_pnl=qty*(exit-entry)", "stock_ret_%"]
    for c, h in enumerate(h4, 1):
        cell = ws4.cell(1, c, h); cell.font = H; cell.fill = HDR; cell.border = BORD; cell.alignment = Cc
    ri = 2
    for x in dc:
        if not x["traded"]:
            continue
        for s in x["stocks"]:
            vals = [x["date"], s["rank"], s["symbol"], "09:15", s["entry_price"], s["alloc"], s["qty"],
                    s["deployed"], x["exit_time"], s["exit_price"], x["exit_reason"], s["gross_pnl"], s["stock_ret"]]
            for c, v in enumerate(vals, 1):
                cell = ws4.cell(ri, c, v); cell.font = NORM; cell.alignment = Ll if c in (1, 3, 4, 9, 11) else Rr
                if c in (6, 8, 12): cell.number_format = RS
                if c == 13: cell.number_format = '0.000'
            if s["gross_pnl"] > 0:
                for c in range(1, 14): ws4.cell(ri, c).fill = GRNF
            elif s["gross_pnl"] < 0:
                for c in range(1, 14): ws4.cell(ri, c).fill = REDF
            ri += 1
    for i, w in enumerate([11, 5, 12, 9, 11, 11, 20, 17, 9, 11, 12, 24, 11], 1):
        ws4.column_dimensions[get_column_letter(i)].width = w
    ws4.freeze_panes = "A2"; ws4.auto_filter.ref = f"A1:M{ri}"

    out = ROOT / "docs" / "ops" / "AUTOTRADE_FULL_REPORT_GATED.xlsx"
    wb.save(out)
    print(f"[*] WROTE {out}")


if __name__ == "__main__":
    main()
