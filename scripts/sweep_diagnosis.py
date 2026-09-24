"""Stock-level, day-by-day diagnosis of the context-gated sweep-reversal logic over a
2-month window. For every stock that FIRED the signal (strong-tape -> long swept-low@open;
weak-tape -> short swept-high@prior-day-high; both volume-confirmed), record entry/EOD
outcome (WIN/LOSS) plus diagnostic features, then attribute the LOSERS to 6 reasons each
for long and short. Writes an Excel (day-by-day trade log + reason attribution)."""
import sqlite3
from pathlib import Path
from collections import defaultdict
import numpy as np
import pandas as pd

ROOT = Path(r"C:\Users\SPS\Desktop\Kanida.ai Terminal Quant Intelligence Engine")
DB = ROOT / "universe_engine" / "data" / "db" / "kanida_universe.db"
OUT = ROOT / "docs" / "ops" / "SWEEP_STOCK_DIAGNOSIS.xlsx"
IDX = {"NIFTY", "BANKNIFTY", "FINNIFTY", "MIDCPNIFTY", "NIFTYNXT50"}
WIN = ("2025-09-01", "2025-10-31")          # the random two months
DET = ("09:16", "09:45"); RECLAIM_K = 3; VPM_DIV = 375.0
STRONG, WEAK = 33.0, 66.0; RVOL_MIN = 1.5


def load():
    con = sqlite3.connect(str(DB))
    fo = [r[0] for r in con.execute("SELECT symbol FROM fo_stock_master").fetchall() if r[0] not in IDX]
    ph = ",".join("?" * len(fo))
    daily = pd.read_sql_query(
        f"SELECT symbol, substr(bar_time,1,10) d, min(low) lo, max(high) hi, sum(volume) dvol, "
        f"max(CASE WHEN substr(bar_time,12,5)='15:29' THEN close END) dclose "
        f"FROM ohlc_1min WHERE symbol IN ({ph}) AND substr(bar_time,1,10) BETWEEN '2025-07-15' AND ? "
        f"GROUP BY symbol, substr(bar_time,1,10)", con, params=fo + [WIN[1]])
    bars = pd.read_sql_query(
        f"SELECT symbol, substr(bar_time,1,10) d, substr(bar_time,12,5) hm, open, high, low, close, volume "
        f"FROM ohlc_1min WHERE symbol IN ({ph}) AND substr(bar_time,1,10) BETWEEN ? AND ? ORDER BY symbol, bar_time",
        con, params=fo + [WIN[0], WIN[1]])
    con.close()
    daily = daily.sort_values(["symbol", "d"])
    for c in ["lo", "hi", "dclose", "dvol"]:
        daily["p_" + c] = daily.groupby("symbol")[c].shift(1)
    daily["p_3d"] = daily.groupby("symbol").dclose.pct_change(3).shift(1) * 100   # trailing 3-day move before today
    return fo, daily, bars


def main():
    fo, daily, bars = load()
    lvl = daily.set_index(["symbol", "d"])[["p_lo", "p_hi", "p_dclose", "p_dvol", "dclose", "p_3d"]].to_dict("index")
    O = bars.open.values; H = bars.high.values; Lw = bars.low.values; C = bars.close.values; V = bars.volume.values
    HM = bars.hm.values.astype("U5")
    idxmap = bars.groupby(["symbol", "d"], sort=False).indices

    # ---- pass 1: breadth + market EOD per day ----
    early = defaultdict(list); mkt = defaultdict(list)
    grp = {}
    for (sym, d), pos in idxmap.items():
        grp[(sym, d)] = pos
        hm = HM[pos]; c = C[pos]; o = O[pos]
        if len(pos) < 20 or hm[0] != "09:15":
            continue
        i918 = np.where(hm == "09:18")[0]
        if len(i918):
            early[d].append(c[i918[0]] / o[0] - 1)
        Lm = lvl.get((sym, d))
        if Lm and Lm.get("dclose"):
            mkt[d].append(Lm["dclose"] / o[0] - 1)
    breadth = {d: np.mean(np.array(x) < 0) * 100 for d, x in early.items()}
    mkt_eod = {d: np.mean(x) * 100 for d, x in mkt.items()}

    # ---- pass 2: gated signal detection with full detail ----
    rows = []
    for (sym, d), pos in grp.items():
        bd = breadth.get(d)
        if bd is None:
            continue
        strong = bd <= STRONG; weak = bd >= WEAK
        if not (strong or weak):
            continue
        hm = HM[pos]; o = O[pos]; h = H[pos]; l = Lw[pos]; c = C[pos]; v = V[pos]
        if len(pos) < 20 or hm[0] != "09:15":
            continue
        Lm = lvl.get((sym, d))
        if not Lm or not Lm.get("p_dvol") or Lm["p_dvol"] <= 0:
            continue
        opn = o[0]; rvol = v / (Lm["p_dvol"] / VPM_DIV); dayclose = Lm["dclose"]
        n = len(c); det = np.where((hm >= DET[0]) & (hm <= DET[1]))[0]
        if not len(det) or not dayclose:
            continue
        gap = (opn / Lm["p_dclose"] - 1) * 100 if Lm.get("p_dclose") else np.nan

        side, lvlname, Lv = (("long", "open", opn) if strong else ("short", "pd_hi", Lm.get("p_hi")))
        if Lv is None or not np.isfinite(Lv):
            continue
        sw = det[(l[det] < Lv)] if side == "long" else det[(h[det] > Lv)]
        if not len(sw):
            continue
        i = int(sw[0]); recl = None
        for j in range(i, min(i + RECLAIM_K + 1, n)):
            if (side == "long" and c[j] > Lv) or (side == "short" and c[j] < Lv):
                recl = j; break
        if recl is None or not (rvol[recl] > RVOL_MIN):     # gated: reclaimed + volume-confirmed
            continue
        entry = c[recl]; sgn = 1 if side == "long" else -1
        eod_ret = (dayclose / entry - 1) * sgn * 100
        path = c[recl:]                                     # entry -> EOD path (within pulled day)
        mfe = ((path.max() if side == "long" else path.min()) / entry - 1) * sgn * 100
        mae = ((path.min() if side == "long" else path.max()) / entry - 1) * sgn * 100
        after = c[recl + 1:]
        re_break = bool(np.any(after < Lv)) if side == "long" else bool(np.any(after > Lv))
        quick_fail = bool(np.any(after[:10] < Lv)) if side == "long" else bool(np.any(after[:10] > Lv))
        rows.append(dict(date=d, side=side, tape=("strong" if strong else "weak"), symbol=sym, level=lvlname,
                         entry=round(float(entry), 2), eod=round(float(dayclose), 2), ret=round(float(eod_ret), 2),
                         win=eod_ret > 0, gap=round(float(gap), 2), prior3d=round(float(Lm.get("p_3d") or np.nan), 2),
                         reclaim_rvol=round(float(rvol[recl]), 2), re_break=re_break, quick_fail=quick_fail,
                         mfe=round(float(mfe), 2), mae=round(float(mae), 2),
                         mkt_eod=round(float(mkt_eod.get(d, np.nan)), 3)))
    df = pd.DataFrame(rows).sort_values(["date", "side", "symbol"])
    if df.empty:
        print("no gated trades in window"); return

    # ---- reason attribution for losers ----
    def reasons(r):
        out = []
        if r.side == "long":
            if r.mfe >= 0.5 and r.ret < 0: out.append("1_faded_into_close")   # worked intraday then gave it back (EXIT)
            if r.mfe < 0.3: out.append("2_no_thrust")                          # reclaim had no follow-through (weak reversal)
            if r.quick_fail: out.append("3_quick_fail")                        # re-broke the low within 10 min (false reclaim)
            if r.mkt_eod < 0: out.append("4_tape_reversed")                    # strong-open tape faded to a red close (CONTEXT)
            if r.gap < -1.5: out.append("5_news_gapdown")                      # gapped down hard = genuine selling
            if r.prior3d < -4: out.append("6_own_downtrend")                   # fighting the stock's own downtrend
        else:
            if r.mfe >= 0.5 and r.ret < 0: out.append("1_faded_into_close")
            if r.mfe < 0.3: out.append("2_no_thrust")
            if r.quick_fail: out.append("3_quick_fail")
            if r.mkt_eod > 0: out.append("4_tape_reversed")                    # weak-open tape squeezed to a green close
            if r.gap > 1.5: out.append("5_news_gapup")
            if r.prior3d > 4: out.append("6_own_uptrend")
        return out
    df["reasons"] = df.apply(lambda r: ",".join(reasons(r)) if not r.win else "", axis=1)

    # ---- print summary ----
    print(f"Window {WIN[0]}..{WIN[1]} | gated trades: {len(df)} "
          f"(long {int((df.side=='long').sum())}, short {int((df.side=='short').sum())})")
    for side in ["long", "short"]:
        s = df[df.side == side]; w = s.win.mean() * 100 if len(s) else 0
        print(f"\n{side.upper()}: {len(s)} trades | win {w:.0f}% | avg {s.ret.mean():+.2f}% | "
              f"winners avg {s[s.win].ret.mean():+.2f}% | losers avg {s[~s.win].ret.mean():+.2f}%")
        los = s[~s.win]
        if len(los):
            print(f"  WHY LOSERS FAILED ({len(los)} losers) — % of losers exhibiting each reason:")
            allr = [rr for _, r in los.iterrows() for rr in (r.reasons.split(",") if r.reasons else [])]
            wr = df[(df.side == side) & df.win]
            for reason, cnt in sorted(pd.Series(allr).value_counts().items(), key=lambda x: -x[1]):
                win_share = (wr.apply(lambda w: reason in reasons(w), axis=1).mean() * 100) if len(wr) else 0
                print(f"    {reason:<20} {cnt:>3}/{len(los)} losers ({cnt/len(los)*100:.0f}%)   [winners w/ same: {win_share:.0f}%]")

    # ---- write Excel ----
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    F = "Calibri"; H = Font(name=F, bold=True, size=9, color="FFFFFF"); HDR = PatternFill("solid", start_color="1F6F8B")
    GRNF = PatternFill("solid", start_color="E7F4EA"); REDF = PatternFill("solid", start_color="FDECEA")
    thin = Side(style="thin", color="DDDDDD"); BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = Workbook(); ws = wb.active; ws.title = "Trade Log (day-by-day)"
    cols = ["date", "side", "tape", "symbol", "entry", "eod", "ret", "win", "gap", "prior3d",
            "reclaim_rvol", "quick_fail", "mfe", "mae", "mkt_eod", "reasons"]
    for c, h in enumerate(cols, 1):
        x = ws.cell(1, c, h); x.font = H; x.fill = HDR; x.border = BORD
    for ri, (_, r) in enumerate(df.iterrows(), 2):
        for c, k in enumerate(cols, 1):
            x = ws.cell(ri, c, r[k]); x.border = BORD; x.font = Font(name=F, size=9)
        fill = GRNF if r.win else REDF
        for c in range(1, len(cols) + 1):
            ws.cell(ri, c).fill = fill
    ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:P{len(df)+1}"
    for i, wdt in enumerate([11, 6, 7, 12, 9, 9, 7, 6, 7, 8, 12, 9, 7, 7, 8, 34], 1):
        ws.column_dimensions[chr(64 + i)].width = wdt

    # reason attribution sheet
    ws2 = wb.create_sheet("Loser Reasons")
    ws2.cell(1, 1, "Why losers went against us — % of losing trades exhibiting each reason").font = Font(name=F, bold=True, size=12, color="1F6F8B")
    r = 3
    for side in ["long", "short"]:
        los = df[(df.side == side) & (~df.win)]
        ws2.cell(r, 1, f"{side.upper()} losers: {len(los)}").font = Font(name=F, bold=True, size=11); r += 1
        for c, h in enumerate(["reason", "# losers", "% losers"], 1):
            x = ws2.cell(r, c, h); x.font = H; x.fill = HDR
        r += 1
        allr = [rr for _, rr2 in los.iterrows() for rr in (rr2.reasons.split(",") if rr2.reasons else [])]
        for reason, cnt in sorted(pd.Series(allr).value_counts().items(), key=lambda x: -x[1]) if len(allr) else []:
            ws2.cell(r, 1, reason); ws2.cell(r, 2, cnt); ws2.cell(r, 3, f"{cnt/len(los)*100:.0f}%"); r += 1
        r += 1
    for i, wdt in enumerate([22, 10, 10], 1):
        ws2.column_dimensions[chr(64 + i)].width = wdt

    try:
        wb.save(OUT); out = OUT
    except PermissionError:
        out = OUT.with_name(OUT.stem + "_v2.xlsx"); wb.save(out)
    print(f"\n[*] WROTE {out}")


if __name__ == "__main__":
    main()
