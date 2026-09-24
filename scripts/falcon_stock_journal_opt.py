"""Per-stock trade journal for EVERY day under the OPTIMIZED config
(arm 2.5 / floor 1 / giveback 1.5 / stop 3.0, cash Top-5, next-open fills).
Includes extended days to 2026-07-03. Writes a NEW xlsx + pickles the full
per-day dataset (_opt_dataset.pkl) so the regime filter reuses it (no re-fetch).
"""
import sys, pickle, sqlite3
sys.path.insert(0, r"C:\Users\SPS\Desktop\Kanida.ai Terminal Quant Intelligence Engine\scripts")
from pathlib import Path
from collections import defaultdict
import numpy as np
from falcon_optimized_monthly import (ROOT, RND, rnd_signals, extended_signals,
                                      rnd_1min, kite_1min, build, RND_1MIN_MAX)

ARM, FLOOR, GIVE, STOP = 2.5, 1.0, 1.5, 3.0


def assemble_full():
    con = sqlite3.connect(str(RND)); base = rnd_signals(con); ext = extended_signals()
    from services.kite_auth import get_kite_client
    kite = get_kite_client(check=False)
    tok = {i["tradingsymbol"]: i["instrument_token"] for i in kite.instruments("NSE")}
    mats = []
    for d in sorted(base):
        M = build(rnd_1min(con, d, base[d]), base[d])
        if M:
            mats.append((d, M))
    for d in sorted(ext):
        per = rnd_1min(con, d, ext[d]) if d <= RND_1MIN_MAX else kite_1min(kite, tok, d, ext[d])
        M = build(per, ext[d])
        if M:
            mats.append((d, M))
    con.close(); mats.sort(key=lambda x: x[0])
    return mats


def basket_exit(M):
    """Return (exit_bar_index, reason) for the optimized basket trail (next-open fill)."""
    qty, close, dep, n = M["qty"], M["close"], M["dep"], M["n"]
    ret = (close @ qty - dep) / dep * 100.0
    armed = False; peak = None
    for i in range(n - 1):
        r = ret[i]
        if r <= -STOP:
            return i, "STOP"
        if not armed:
            if r >= ARM:
                armed = True; peak = r
            continue
        peak = max(peak, r)
        if r <= max(FLOOR, peak - GIVE):
            return i, ("FLOOR" if max(FLOOR, peak - GIVE) == FLOOR else "TRAIL")
    return n - 1, "EOD"


def main():
    mats = assemble_full()
    pickle.dump(mats, open(ROOT / "docs" / "ops" / "_opt_dataset.pkl", "wb"))
    print(f"[*] assembled + pickled {len(mats)} days ({mats[0][0]} -> {mats[-1][0]})", flush=True)

    stock_rows = []; day_rows = []
    for d, M in mats:
        xb, reason = basket_exit(M)
        n = M["n"]; entry = M["entry"]; qty = M["qty"]; close = M["close"]; high = M["high"]; low = M["low"]; opn = M["opn"]
        grid = M["grid"]; dep = M["dep"]
        # portfolio realized (next-open fill on trigger; close on EOD)
        if reason == "EOD":
            exit_px = close[n - 1]
        else:
            exit_px = opn[xb + 1]
            # guard: if next-open missing/<=0 use close at xb
            exit_px = np.where(np.isfinite(exit_px) & (exit_px > 0), exit_px, close[xb])
        port_ret = float(((exit_px - entry) * qty).sum() / dep * 100.0)
        exit_time = grid[min(xb + 1 if reason != "EOD" else n - 1, n - 1)]
        for j in range(M["nstocks"]):
            e = entry[j]; xp = float(exit_px[j])
            mae = float((low[0:xb + 1, j].min() - e) / e * 100)
            mfe = float((high[:, j].max() - e) / e * 100)
            posthi = float((high[xb + 1:, j].max() - e) / e * 100) if xb + 1 < n else (xp / e - 1) * 100
            eod = float((close[n - 1, j] - e) / e * 100)
            stock_rows.append([d, d[:7], M["rank"][j], M["syms"][j], round(e, 2), int(qty[j]),
                               round(xp, 2), exit_time, reason, round((xp / e - 1) * 100, 3),
                               round((xp - e) * qty[j], 1), round(mae, 3), round(mfe, 3),
                               round(posthi, 3), round(eod, 3), round(port_ret, 3)])
        day_rows.append([d, d[:7], M["nstocks"], ", ".join(M["syms"]), exit_time, reason,
                         round(port_ret, 3), "WIN" if port_ret > 0 else "LOSS"])

    # ---- write xlsx ----
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    F = "Calibri"
    TITLE = Font(name=F, bold=True, size=13, color="1F6F8B"); NOTE = Font(name=F, size=9, italic=True, color="666666")
    H = Font(name=F, bold=True, size=9, color="FFFFFF"); HDR = PatternFill("solid", start_color="1F6F8B")
    NORM = Font(name=F, size=10); GRNF = PatternFill("solid", start_color="E7F4EA"); REDF = PatternFill("solid", start_color="FDECEA")
    thin = Side(style="thin", color="DDDDDD"); BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
    Cc = Alignment(horizontal="center"); Rr = Alignment(horizontal="right"); Ll = Alignment(horizontal="left")
    wb = Workbook()
    # sheet 1: per-stock journal
    ws = wb.active; ws.title = "Stock Journal"
    hs = ["date", "month", "rank", "symbol", "entry", "qty", "exit_px", "exit_time", "exit_reason",
          "stock_ret%", "pnl_rs", "MAE%", "MFE%(day)", "post_exit_hi%", "EOD_close%", "PORT_ret%"]
    ws.cell(1, 1, "Per-Stock Trade Journal — OPTIMIZED config (arm 2.5 / floor 1 / giveback 1.5 / stop 3.0)").font = TITLE
    ws.cell(2, 1, "Cash Top-5, entry 09:15, basket-level exit (all names exit together on the trail/stop/EOD trigger; next-open fills). 530 days incl. to 2026-07-03.").font = NOTE
    for c, h in enumerate(hs, 1):
        cell = ws.cell(4, c, h); cell.font = H; cell.fill = HDR; cell.border = BORD; cell.alignment = Cc
    r = 5
    for row in stock_rows:
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v); cell.font = NORM; cell.alignment = (Ll if c in (1, 2, 4, 8, 9) else Rr)
            if c in (10, 12, 13, 14, 15, 16): cell.number_format = '0.000'
        # tint by stock result
        if row[9] > 0:
            for c in range(1, 17): ws.cell(r, c).fill = GRNF
        elif row[9] < 0:
            for c in range(1, 17): ws.cell(r, c).fill = REDF
        r += 1
    for i, w in enumerate([11, 8, 5, 12, 9, 6, 9, 9, 10, 10, 10, 8, 10, 12, 11, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"; ws.auto_filter.ref = f"A4:P{len(stock_rows)+4}"
    # sheet 2: daily summary
    ws2 = wb.create_sheet("Daily Summary")
    hd = ["date", "month", "n_stocks", "stocks", "exit_time", "exit_reason", "PORT_ret%", "result"]
    for c, h in enumerate(hd, 1):
        cell = ws2.cell(1, c, h); cell.font = H; cell.fill = HDR; cell.border = BORD; cell.alignment = Cc
    for ri, row in enumerate(day_rows, 2):
        for c, v in enumerate(row, 1):
            cell = ws2.cell(ri, c, v); cell.font = NORM; cell.alignment = (Ll if c in (1, 2, 4, 5, 6, 8) else Rr)
            if c == 7: cell.number_format = '0.000'
        if row[6] > 0:
            for c in range(1, 9): ws2.cell(ri, c).fill = GRNF
        elif row[6] < 0:
            for c in range(1, 9): ws2.cell(ri, c).fill = REDF
    for i, w in enumerate([11, 8, 8, 46, 9, 10, 10, 8], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A2"; ws2.auto_filter.ref = f"A1:H{len(day_rows)+1}"
    out = ROOT / "docs" / "ops" / "AUTOTRADE_STOCK_JOURNAL_OPTIMIZED.xlsx"
    wb.save(out)
    print(f"[*] WROTE {out}  ({len(stock_rows)} stock-day rows, {len(day_rows)} days)")
    reasons = defaultdict(int)
    for row in day_rows:
        reasons[row[5]] += 1
    print("  exit-reason mix (days):", dict(reasons))


if __name__ == "__main__":
    main()
