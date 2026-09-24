"""Build the full walk-forward review workbook from the generated CSVs + grid JSON."""
import csv, json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]; OPS = ROOT / "docs" / "ops"
F = "Calibri"
TITLE = Font(name=F, bold=True, size=14, color="1F6F8B"); SUB = Font(name=F, bold=True, size=11, color="222222")
H = Font(name=F, bold=True, size=9, color="FFFFFF"); HDR = PatternFill("solid", start_color="1F6F8B")
BOLD = Font(name=F, bold=True, size=10); NORM = Font(name=F, size=10); NOTE = Font(name=F, size=9, italic=True, color="666666")
GRN = Font(name=F, bold=True, size=10, color="1E7E34"); RED = Font(name=F, bold=True, size=10, color="B00020")
GRNF = PatternFill("solid", start_color="E7F4EA"); REDF = PatternFill("solid", start_color="FDECEA"); AMBF = PatternFill("solid", start_color="FFF4CE")
thin = Side(style="thin", color="DDDDDD"); BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
Cc = Alignment(horizontal="center"); Rr = Alignment(horizontal="right"); Ll = Alignment(horizontal="left", wrap_text=True)


def readcsv(name):
    with open(OPS / name) as f:
        return list(csv.reader(f))


day = readcsv("wf_day_analysis.csv")
stock = readcsv("wf_stock_journal.csv")
month = readcsv("wf_monthly_consistency.csv")
grid = json.loads((OPS / "param_grid.json").read_text())

wb = Workbook()

# ── Overview ──
ws = wb.active; ws.title = "Overview"
for i, w in enumerate([3, 34, 16, 16, 40], 1): ws.column_dimensions[get_column_letter(i)].width = w
r = 2
ws.cell(r, 2, "Walk-Forward Trade-Quality & Parameter Review").font = TITLE; r += 1
ws.cell(r, 2, "517 trading days · Top-5 · CASH (non-MTF) portfolio returns · true walk-forward, no leakage, all days.").font = NOTE; r += 2
allr = [float(x[3]) for x in day[1:]]
import statistics as st
ge1 = sum(1 for v in allr if v >= 1) / len(allr) * 100
pos = sum(1 for v in allr if v > 0) / len(allr) * 100
ws.cell(r, 2, "YOUR TARGET vs REALITY").font = SUB; r += 1
for lab, val, tone in [
    ("Target: avg daily return (cash)", ">= +1.00%", None),
    ("  Achieved (current params)", f"+{st.mean(allr):.3f}%  ✓ (mean)", "1E7E34"),
    ("Target: >= +1% on 90% of days each month", "90% of days", None),
    ("  Achieved (current params)", f"{ge1:.1f}% of days · 0 of 26 months pass", "B00020"),
    ("  Achieved (BEST of 100 param combos)", f"{grid[0]['full']['ge1']:.1f}% of days · 0 of 26 months pass", "B00020"),
    ("Positive days (current)", f"{pos:.1f}%", None),
]:
    ws.cell(r, 2, lab).font = BOLD if not lab.startswith("  ") else NORM
    c = ws.cell(r, 3, val); c.alignment = Rr
    c.font = Font(name=F, bold=True, size=10, color=tone) if tone else NORM
    r += 1
r += 1
ws.cell(r, 2, "VERDICT: the '≥1% on 90% of days' target is NOT reachable — best of 100 combos hits 61.7% of days; 0/26 months pass.").font = RED
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5); r += 1
ws.cell(r, 2, "It requires the 10th-percentile day to be ≥+1% (i.e. almost no flat/down days) — impossible with real market dispersion.").font = NOTE
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5); r += 2
ws.cell(r, 2, "BEST-CONSISTENCY combos (ranked by % days ≥ +1%, cash):").font = SUB; r += 1
gh = ["stop%", "arm", "give%", "floor%", "mean%", "med%", "days≥1%", "pos%", "worst%", "test days≥1%"]
for c, h in enumerate(gh, 1):
    cell = ws.cell(r, c + 1, h); cell.font = H; cell.fill = HDR; cell.border = BORD; cell.alignment = Cc
r += 1
for g in grid[:6]:
    f = g["full"]; arm = "none" if g["arm"] >= 9 else f"{g['arm']*100:.1f}"
    vals = [g["stop"]*100, arm, g["give"]*100, g["floor"]*100, f["mean"], f["median"], f["ge1"], f["pos"], f["worst"], g["test"]["ge1"]]
    for c, v in enumerate(vals, 2):
        cell = ws.cell(r, c, v); cell.border = BORD; cell.alignment = Rr; cell.font = NORM
        if isinstance(v, float): cell.number_format = '0.00'
    for c in range(2, 12): ws.cell(r, c).fill = GRNF
    r += 1
# current row
cur = next((g for g in grid if abs(g['stop']-0.015) < 1e-9 and abs(g['arm']-0.02) < 1e-9 and abs(g['give']-0.005) < 1e-9 and abs(g['floor']-0.01) < 1e-9), None)
if cur:
    f = cur["full"]; vals = [1.5, "2.0", 0.5, 1.0, f["mean"], f["median"], f["ge1"], f["pos"], f["worst"], cur["test"]["ge1"]]
    for c, v in enumerate(vals, 2):
        cell = ws.cell(r, c, v); cell.border = BORD; cell.alignment = Rr; cell.font = NORM
        if isinstance(v, float): cell.number_format = '0.00'
        cell.fill = AMBF
    ws.cell(r, 12, "  ← CURRENT live").font = NOTE

# ── Per-Day (517) ──
def datasheet(name, header, rows, colw, pct_cols=(), flag_col=None):
    ws = wb.create_sheet(name)
    for i, w in enumerate(colw, 1): ws.column_dimensions[get_column_letter(i)].width = w
    for c, h in enumerate(header, 1):
        cell = ws.cell(1, c, h); cell.font = H; cell.fill = HDR; cell.border = BORD; cell.alignment = Cc
    for ri, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            try: v2 = float(v)
            except (ValueError, TypeError): v2 = v
            cell = ws.cell(ri, c, v2); cell.font = NORM
            cell.alignment = Rr if isinstance(v2, float) else Ll
            if isinstance(v2, float) and c in pct_cols: cell.number_format = '0.000'
        if flag_col is not None and str(row[flag_col-1]).upper() in ("YES", "PASS"):
            for c in range(1, len(header)+1): ws.cell(ri, c).fill = GRNF
        elif flag_col is not None and str(row[flag_col-1]).upper() in ("NO", "FAIL"):
            for c in range(1, len(header)+1): ws.cell(ri, c).fill = REDF
    ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(rows)+1}"


datasheet("Per-Day (517d)", day[0], day[1:], [11, 8, 8, 14, 13, 11, 11, 11, 7, 9, 9, 16], pct_cols=(4, 6, 7, 8), flag_col=11)
datasheet("Daily Journal (per-stock)", stock[0], stock[1:],
          [11, 8, 5, 12, 9, 6, 9, 9, 12, 9, 9, 8, 8, 8, 8, 11, 11], pct_cols=(10, 12, 14, 16, 17))
datasheet("Monthly Consistency", month[0], month[1:], [9, 8, 12, 12, 12, 13, 11, 14], pct_cols=(3, 4, 5, 6), flag_col=8)

# ── Param Grid (all) ──
ws = wb.create_sheet("Param Grid (100)")
gh = ["stop%", "arm%", "give%", "floor%", "mean%", "median%", "days≥1%", "pos%", "worst%", "months_pass", "TRAIN days≥1%", "TEST days≥1%"]
for c, h in enumerate(gh, 1):
    cell = ws.cell(1, c, h); cell.font = H; cell.fill = HDR; cell.border = BORD; cell.alignment = Cc
for ri, g in enumerate(grid, 2):
    f = g["full"]; arm = 0 if g["arm"] >= 9 else g["arm"]*100
    vals = [g["stop"]*100, arm, g["give"]*100, g["floor"]*100, f["mean"], f["median"], f["ge1"], f["pos"], f["worst"],
            f["months_pass"], g["train"]["ge1"], g["test"]["ge1"]]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(ri, c, v); cell.font = NORM; cell.alignment = Rr; cell.number_format = '0.00'
for i, w in enumerate([7, 7, 7, 7, 8, 8, 8, 7, 8, 11, 12, 11], 1): ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:L{len(grid)+1}"

out = OPS / "AUTOTRADE_WALKFORWARD_ANALYSIS_517d.xlsx"
wb.save(out)
print("WROTE", out, "| sheets:", wb.sheetnames)
