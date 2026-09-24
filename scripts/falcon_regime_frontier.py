"""Regime-filter efficiency frontier: for each filter, decompose the SKIPPED days
into gains-forgone (skipped positives) vs losses-avoided (skipped negatives), and
rank by return-to-drawdown (Calmar-like). Finds the true 'sweet spot' middle path.
Reuses falcon_regime_filter (same pickle, same Nifty series). Cash, walk-forward.
"""
import sys
sys.path.insert(0, r"C:\Users\SPS\Desktop\Kanida.ai Terminal Quant Intelligence Engine\scripts")
sys.path.insert(0, r"C:\Users\SPS\Desktop\Kanida.ai Terminal Quant Intelligence Engine\backend")
from collections import defaultdict
import numpy as np
from falcon_regime_filter import ROOT, strategy_returns, proxy_series, regimes, selfreg


def evaluate(dates, rets, traded):
    tr = np.array([rets[d] for d in dates if traded[d]])
    alln = np.array([rets[d] if traded[d] else 0.0 for d in dates])
    skipped = [rets[d] for d in dates if not traded[d]]
    sp = [x for x in skipped if x > 0]; sn = [x for x in skipped if x < 0]
    eq = np.cumsum(alln); dd = float((np.maximum.accumulate(eq) - eq).max())
    total = float(tr.sum())
    return dict(
        traded=len(tr), skip=len(skipped),
        mean_tr=round(float(tr.mean()), 3), ge1=round(float((tr >= 1).mean() * 100), 1),
        pos=round(float((tr > 0).mean() * 100), 1), worst=round(float(tr.min()), 2),
        total=round(total, 1), maxdd=round(dd, 1),
        ret_dd=round(total / dd, 1) if dd > 0 else None,
        skip_pos_n=len(sp), skip_pos_sum=round(sum(sp), 1),
        skip_neg_n=len(sn), skip_neg_sum=round(sum(sn), 1),
        eff=round(abs(sum(sn)) / sum(sp), 3) if sp else None)  # loss-avoided per gain-forgone


def main():
    dates, rets = strategy_returns()
    ser = proxy_series(dates)
    reg = regimes(dates, ser); s10 = selfreg(dates, rets, 10); s5 = selfreg(dates, rets, 5)

    FILTERS = {
        "BASE (trade all)": lambda d: True,
        "Nifty>10D SMA": lambda d: reg[d].get("above10", reg[d]["above20"]),
        "Nifty>20D SMA": lambda d: reg[d]["above20"],
        "Nifty>50D SMA": lambda d: reg[d]["above50"],
        "Nifty 5dret>-1%": lambda d: reg[d].get("ret5", 0) > -1.0,
        "Nifty 5dret>-1.5%": lambda d: reg[d].get("ret5", 0) > -1.5,
        "Nifty 5dret>-2%": lambda d: reg[d].get("ret5", 0) > -2.0,
        "Nifty 5dret>-3%": lambda d: reg[d].get("ret5", 0) > -3.0,
        "Self trail5>0": lambda d: s5[d] > 0,
        "Self trail10>0": lambda d: s10[d] > 0,
        "COMBO 20D & self>0": lambda d: reg[d]["above20"] and s10[d] > 0,
        "COMBO 5dret>-2% & self>0": lambda d: reg[d].get("ret5", 0) > -2.0 and s10[d] > 0,
        "COMBO 5dret>-2% & self5>0": lambda d: reg[d].get("ret5", 0) > -2.0 and s5[d] > 0,
    }
    res = {name: evaluate(dates, rets, {d: bool(fn(d)) for d in dates}) for name, fn in FILTERS.items()}

    print(f"{'filter':<27}{'trade':>6}{'total%':>8}{'maxDD':>7}{'ret/DD':>7}{'pos%':>6}{'d>=1%':>6}{'worst':>7}{'forgone':>9}{'avoided':>9}{'eff':>6}")
    for n, m in res.items():
        print(f"{n:<27}{m['traded']:>6}{m['total']:>8.0f}{m['maxdd']:>7.1f}{str(m['ret_dd']):>7}{m['pos']:>6}{m['ge1']:>6}{m['worst']:>7}"
              f"{m['skip_pos_sum']:>9.0f}{m['skip_neg_sum']:>9.0f}{str(m['eff']):>6}")

    # verify user's COMBO decomposition
    c = res["COMBO 20D & self>0"]
    print(f"\n[verify] COMBO skipped {c['skip']} days = {c['skip_pos_n']} pos (+{c['skip_pos_sum']}%) + "
          f"{c['skip_neg_n']} neg ({c['skip_neg_sum']}%). Net forgone {c['skip_pos_sum']+c['skip_neg_sum']:.0f}%.")

    # rank by ret/DD
    ranked = sorted([(m['ret_dd'], n) for n, m in res.items() if m['ret_dd']], reverse=True)
    print("\nBEST risk-adjusted (return / maxDD):")
    for rd, n in ranked[:5]:
        m = res[n]
        print(f"  {n:<27} ret/DD {rd:>6} | total +{m['total']:.0f}% | DD {m['maxdd']}% | pos {m['pos']}% | d>=1% {m['ge1']}%")

    # write sheet
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    path = ROOT / "docs" / "ops" / "AUTOTRADE_REGIME_FILTER.xlsx"
    wb = load_workbook(path)
    if "Frontier" in wb.sheetnames:
        del wb["Frontier"]
    ws = wb.create_sheet("Frontier")
    F = "Calibri"; H = Font(name=F, bold=True, size=9, color="FFFFFF"); HDR = PatternFill("solid", start_color="1F6F8B")
    NORM = Font(name=F, size=10); TITLE = Font(name=F, bold=True, size=13, color="1F6F8B"); NOTE = Font(name=F, size=9, italic=True, color="666666")
    GRNF = PatternFill("solid", start_color="E7F4EA"); TOT = PatternFill("solid", start_color="FFF4CE")
    thin = Side(style="thin", color="DDDDDD"); BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
    Cc = Alignment(horizontal="center"); Rr = Alignment(horizontal="right"); Ll = Alignment(horizontal="left")
    ws.cell(1, 1, "Regime-Filter Efficiency Frontier").font = TITLE
    ws.cell(2, 1, "forgone = total return on SKIPPED positive days (opportunity cost). avoided = total on skipped negatives. "
                  "eff = |avoided|/forgone (higher=skips mostly losers). ret/DD = total return / max drawdown (Calmar-like).").font = NOTE
    heads = ["Filter", "Traded", "Total%", "MaxDD%", "Ret/DD", "Pos%", "Days>=1%", "Worst%",
             "Gains forgone%", "Losses avoided%", "Skip eff"]
    for c, h in enumerate(heads, 1):
        cell = ws.cell(4, c, h); cell.font = H; cell.fill = HDR; cell.border = BORD; cell.alignment = Cc
    r = 5
    best_rd = max(m['ret_dd'] for m in res.values() if m['ret_dd'])
    for n, m in res.items():
        vals = [n, m['traded'], m['total'], m['maxdd'], m['ret_dd'], m['pos'], m['ge1'], m['worst'],
                m['skip_pos_sum'], m['skip_neg_sum'], m['eff']]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v); cell.font = NORM; cell.alignment = Ll if c == 1 else Rr; cell.border = BORD
        if n.startswith("BASE"):
            for c in range(1, 12): ws.cell(r, c).fill = TOT
        elif m['ret_dd'] == best_rd:
            for c in range(1, 12): ws.cell(r, c).fill = GRNF
        r += 1
    for i, w in enumerate([27, 8, 9, 9, 8, 7, 9, 8, 15, 16, 9], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    wb.save(path)
    print(f"\n[*] added 'Frontier' sheet to {path}")


if __name__ == "__main__":
    main()
