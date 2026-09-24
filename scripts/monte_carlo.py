"""Monte Carlo robustness test on the EXISTING net-return logs of both strategies.
NO new price simulation — we only resample the realized trade/daily returns.

Two independent tests per strategy (10,000 iterations each, net-of-cost returns):
  1) SEQUENCE RESHUFFLE (order luck): permute the return order, rebuild the equity curve,
     record max drawdown. Total return is invariant to order, so this isolates: was the
     observed drawdown a lucky ordering, or typical of these returns?
  2) BOOTSTRAP RESAMPLE (selection luck): resample the returns WITH replacement (same n),
     rebuild total return + max DD. This shows the range of outcomes these trades could
     have produced, and whether the edge survives (5th percentile still > 0?).

Return unit = one day's return on a fixed Rs5,00,000 account:
  Intraday   = daily net% (one Rs5L round-trip/day).
  Positional = per-basket net% / 3 (each basket = Rs1.67L sleeve of the 3-sleeve Rs5L book),
               ~1 basket/day, so both series are ~daily returns on the same Rs5L account.

SCOPE (kept deliberately distinct):
  - Monte Carlo here = SEQUENCE + SELECTION robustness of the REALIZED trades.
  - It does NOT model unseen market regimes  -> walk-forward validation covers that.
  - It does NOT model live slippage/impact    -> the market-impact study covers that.
"""
import numpy as np
from datetime import date
from pathlib import Path
import pos_sim as P
from pos_capital import positional_trades, intraday_trades

ROOT = Path(r"C:\Users\SPS\Desktop\Kanida.ai Terminal Quant Intelligence Engine")
OUT = ROOT / "docs" / "ops" / "FALCON_MONTECARLO.xlsx"
NITER = 10000
SEED = 42
MONTH_LEN = 21


def maxdd_rows(mat):
    """Vectorized max drawdown of each row's equity curve (cumsum of returns)."""
    cum = np.cumsum(mat, axis=1)
    peak = np.maximum.accumulate(cum, axis=1)
    return (peak - cum).max(axis=1)


def analyze(r, dates, years, n_iter=NITER, seed=SEED):
    from collections import defaultdict
    rng = np.random.default_rng(seed)
    n = len(r)
    # actual monthly returns (preserve within-month clustering — honest losing-month rate)
    md = defaultdict(float)
    for ri, d in zip(r, dates):
        md[d[:7]] += float(ri)
    months = np.array(sorted(md.values()))
    n_months = len(months); n_lose = int((months < 0).sum())
    worst_month = float(months.min())
    obs_total = float(r.sum())
    obs_dd = float(maxdd_rows(r[None, :])[0])
    obs_annual = obs_total / years

    # ---- BOOTSTRAP (selection luck): resample n WITH replacement ----
    bi = rng.integers(0, n, size=(n_iter, n))
    bsamp = r[bi]
    boot_tot = bsamp.sum(axis=1)
    boot_dd = maxdd_rows(bsamp)
    tot_pcts = np.percentile(boot_tot, [5, 25, 50, 75, 95])
    annual_ci = np.percentile(boot_tot / years, [2.5, 97.5])
    obs_tot_pctile = float((boot_tot < obs_total).mean() * 100)
    p_total_le0 = float((boot_tot <= 0).mean() * 100)

    # ---- SHUFFLE (order luck): permute, record max DD ----
    si = np.argsort(rng.random((n_iter, n)), axis=1)
    ssamp = r[si]
    shuf_dd = maxdd_rows(ssamp)
    dd_pcts = np.percentile(shuf_dd, [5, 25, 50, 75, 95])
    obs_dd_pctile = float((shuf_dd < obs_dd).mean() * 100)

    # ---- Probability of a losing month: EMPIRICAL from actual months (preserves clustering;
    #      a day-bootstrap would wrongly report ~0% because it destroys serial correlation) ----
    p_lose_month = float(n_lose / n_months * 100)

    return dict(
        n=n, years=round(years, 2), obs_total=round(obs_total, 1), obs_annual=round(obs_annual, 1),
        obs_dd=round(obs_dd, 1),
        boot_tot_pctiles={k: round(float(v), 1) for k, v in zip([5, 25, 50, 75, 95], tot_pcts)},
        boot_annual_pctiles={k: round(float(v) / years, 1) for k, v in zip([5, 25, 50, 75, 95], tot_pcts)},
        obs_tot_pctile=round(obs_tot_pctile, 1), p_total_le0=round(p_total_le0, 2),
        annual_ci=[round(float(annual_ci[0]), 1), round(float(annual_ci[1]), 1)],
        shuf_dd_pctiles={k: round(float(v), 1) for k, v in zip([5, 25, 50, 75, 95], dd_pcts)},
        obs_dd_pctile=round(obs_dd_pctile, 1), p_lose_month=round(p_lose_month, 1),
        n_months=n_months, n_lose=n_lose, worst_month=round(worst_month, 1))


def verdict(a, name):
    lines = []
    boot5 = a["boot_tot_pctiles"][5]
    robust_edge = boot5 > 0 and a["p_total_le0"] < 1
    lines.append(
        f"SELECTION (bootstrap): even the 5th-percentile total return is {boot5:+.0f}% "
        f"({'well above 0 -> edge survives resampling' if boot5 > 0 else 'at/below 0 -> FRAGILE'}); "
        f"probability the whole 2-yr result was <=0 is {a['p_total_le0']:.2f}%.")
    dd_med = a["shuf_dd_pctiles"][50]; dd95 = a["shuf_dd_pctiles"][95]
    lucky_dd = a["obs_dd_pctile"] <= 25
    if lucky_dd:
        dd_read = (f"observed drawdown ({a['obs_dd']:.0f}%) is MILDER than {100-a['obs_dd_pctile']:.0f}% of orderings "
                   f"-> we saw a FAVOURABLE sequence. A typical ordering drew down ~{dd_med:.0f}% (95th ~{dd95:.0f}%). "
                   f"Plan for ~{dd_med:.0f}%+ live, not the observed {a['obs_dd']:.0f}%.")
    elif a["obs_dd_pctile"] >= 60:
        dd_read = (f"observed drawdown ({a['obs_dd']:.0f}%) sits at/above the reshuffle median ({dd_med:.0f}%) "
                   f"-> real losses cluster somewhat, so the observed figure is HONEST/conservative, not lucky-low. "
                   f"Size risk to ~{a['obs_dd']:.0f}%.")
    else:
        dd_read = (f"observed drawdown ({a['obs_dd']:.0f}%) is TYPICAL of these returns "
                   f"({a['obs_dd_pctile']:.0f}th pct, median ~{dd_med:.0f}%) -> not a lucky ordering.")
    lines.append("ORDER (reshuffle): " + dd_read)
    lines.append(
        f"Losing months (actual): {a['n_lose']} of {a['n_months']} = {a['p_lose_month']:.0f}%, worst month "
        f"{a['worst_month']:+.0f}%. (Computed on actual months so real loss-clustering is preserved.)")
    lines.append(
        f"Observed total sits at the {a['obs_tot_pctile']:.0f}th percentile of the bootstrap "
        f"(a bootstrap is centred on the observed, so ~50th = a central, typical outcome — not an outlier).")
    lines.append(
        "ASSUMPTION: Monte Carlo treats the realized returns as exchangeable (order/selection). A genuinely "
        "NEW market regime can cluster losses beyond these bands — that is what walk-forward validation covers, "
        "and live slippage is covered by the market-impact study. Keep the three distinct.")
    overall = ("ROBUST — headline return survives both order and selection resampling; observed drawdown is "
               "honest (not a lucky-low ordering)." if robust_edge and not lucky_dd else
               "ROBUST ON RETURN — but the observed DRAWDOWN was a favourable ordering; size risk to the median "
               "reshuffle DD, not the observed." if robust_edge else
               "LUCK-DEPENDENT — the edge does not clearly survive resampling; treat with caution.")
    return overall, lines


# ---------------------------------------------------------------- workbook
def write(intr, posi):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    F = "Calibri"
    TITLE = Font(name=F, bold=True, size=13, color="1F6F8B"); SUB = Font(name=F, bold=True, size=11)
    NOTE = Font(name=F, size=9, italic=True, color="666666"); H = Font(name=F, bold=True, size=9, color="FFFFFF")
    HDR = PatternFill("solid", start_color="1F6F8B"); NORM = Font(name=F, size=10); BOLD = Font(name=F, bold=True, size=10)
    GRN = Font(name=F, size=10, color="1E7E34"); RED = Font(name=F, size=10, color="B00020")
    HILF = PatternFill("solid", start_color="FDF2CC"); GRNF = PatternFill("solid", start_color="E7F4EA")
    thin = Side(style="thin", color="DDDDDD"); BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
    Cc = Alignment(horizontal="center"); Ll = Alignment(horizontal="left", wrap_text=True)
    wb = Workbook()

    def tab(ws, a, name, cfg):
        ov, lines = verdict(a, name)
        ws.cell(1, 1, f"Monte Carlo Robustness — {name}").font = TITLE
        ws.cell(2, 1, cfg).font = NOTE
        ws.cell(3, 1, f"Return series: {a['n']} units over {a['years']} yrs (net of cost). "
                      f"10,000 iterations each for reshuffle and bootstrap.").font = NOTE
        r = 5
        ws.cell(r, 1, "OBSERVED (actual realized sequence)").font = SUB; r += 1
        for lab, v, fmt in [("Total return (2yr)", a["obs_total"], "%"), ("Annualised return", a["obs_annual"], "%"),
                            ("Max drawdown", a["obs_dd"], "%")]:
            ws.cell(r, 1, lab).font = NORM; c = ws.cell(r, 2, f"{v:+.1f}%" if fmt == "%" else v); c.font = BOLD; r += 1
        r += 1
        ws.cell(r, 1, "BOOTSTRAP — total return distribution (selection luck)").font = SUB; r += 1
        hh = ["", "5th", "25th", "50th (median)", "75th", "95th"]
        for c, x in enumerate(hh, 1):
            cell = ws.cell(r, c, x); cell.font = H; cell.fill = HDR; cell.alignment = Cc; cell.border = BORD
        r += 1
        for lab, d in [("Total return %", a["boot_tot_pctiles"]), ("Annualised %", a["boot_annual_pctiles"])]:
            ws.cell(r, 1, lab).font = NORM
            for c, k in enumerate([5, 25, 50, 75, 95], 2):
                x = ws.cell(r, c, f"{d[k]:+.0f}%"); x.font = NORM; x.border = BORD; x.alignment = Cc
                if k == 5: x.font = GRN if d[k] > 0 else RED; x.fill = GRNF if d[k] > 0 else PatternFill("solid", start_color="FDECEA")
            r += 1
        ws.cell(r, 1, "Observed total sits at percentile"); ws.cell(r, 2, f"{a['obs_tot_pctile']:.0f}th").font = BOLD; r += 1
        ws.cell(r, 1, "P(2-yr total <= 0)"); ws.cell(r, 2, f"{a['p_total_le0']:.2f}%").font = BOLD; r += 1
        ws.cell(r, 1, "95% CI, annual return"); ws.cell(r, 2, f"[{a['annual_ci'][0]:+.0f}%, {a['annual_ci'][1]:+.0f}%]").font = BOLD; r += 1
        ws.cell(r, 1, "Losing months (actual)"); ws.cell(r, 2, f"{a['n_lose']}/{a['n_months']} = {a['p_lose_month']:.0f}% (worst {a['worst_month']:+.0f}%)").font = BOLD; r += 2

        ws.cell(r, 1, "RESHUFFLE — max-drawdown distribution (order luck)").font = SUB; r += 1
        for c, x in enumerate(hh, 1):
            cell = ws.cell(r, c, x if x else "sim max DD"); cell.font = H; cell.fill = HDR; cell.alignment = Cc; cell.border = BORD
        r += 1
        ws.cell(r, 1, "Simulated max DD %").font = NORM
        for c, k in enumerate([5, 25, 50, 75, 95], 2):
            x = ws.cell(r, c, f"{a['shuf_dd_pctiles'][k]:.0f}%"); x.font = NORM; x.border = BORD; x.alignment = Cc
        r += 1
        ws.cell(r, 1, "Observed DD vs reshuffle"); ws.cell(r, 2, f"{a['obs_dd']:.0f}% = {a['obs_dd_pctile']:.0f}th pct").font = BOLD; r += 2

        ws.cell(r, 1, "VERDICT").font = SUB; r += 1
        cell = ws.cell(r, 1, ov); cell.font = Font(name=F, bold=True, size=11,
                     color=("1E7E34" if ov.startswith("ROBUST —") else "B8860B" if ov.startswith("ROBUST ON") else "B00020"))
        cell.alignment = Ll; ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6); r += 1
        for ln in lines:
            cell = ws.cell(r, 1, "• " + ln); cell.font = NORM; cell.alignment = Ll
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6); ws.row_dimensions[r].height = 30; r += 1
        for i, w in enumerate([34, 14, 14, 16, 14, 14], 1): ws.column_dimensions[get_column_letter(i)].width = w

    ws1 = wb.active; ws1.title = "Intraday"
    tab(ws1, intr, "Intraday (Validated config)", "Config arm2.5/floor1/give1.5/stop3, Top-5 @09:15, square-off EOD, CASH. Unit = daily net% on Rs5L.")
    ws2 = wb.create_sheet("Positional")
    tab(ws2, posi, "Positional (multi-day)", "Config arm3/floor1/give4/stop6, max-hold 3 sessions, CASH. Unit = per-basket net% on the Rs5L 3-sleeve book (net%/3).")

    ws3 = wb.create_sheet("Method & scope")
    ws3.cell(1, 1, "Method & scope — three distinct robustness questions").font = TITLE
    notes = [
        "This workbook answers ONE question: is the headline return real edge, or one lucky sequence of trades?",
        "It resamples ONLY the realized net-of-cost returns we already have. No new price simulation.",
        "",
        "1) SEQUENCE RESHUFFLE (order luck): permute the return order 10,000x -> distribution of max drawdown.",
        "   Total return is unchanged by order, so this isolates whether the observed DRAWDOWN was lucky ordering.",
        "2) BOOTSTRAP RESAMPLE (selection luck): resample the returns with replacement 10,000x -> distribution of",
        "   total return / drawdown. A bootstrap is centred on the observed, so the robustness signal is the SPREAD:",
        "   if the 5th percentile total is still strongly positive, the edge survives selection luck.",
        "",
        "Kept deliberately DISTINCT (do not conflate):",
        "   - Monte Carlo (here)      = sequence + selection robustness of the REALIZED trades.",
        "   - Walk-forward validation = unseen market regimes (already done).",
        "   - Market-impact study     = live slippage / execution (separate).",
        "",
        "Return unit = one day's return on a fixed Rs5L account. Intraday = daily net%. Positional = per-basket",
        "net%/3 (each basket = Rs1.67L sleeve of the 3-sleeve Rs5L book), ~1 basket/day. Net of ~0.10% round-trip.",
        "Iterations: 10,000 each (reshuffle + bootstrap), seed-fixed for reproducibility.",
    ]
    for i, t in enumerate(notes, 3):
        ws3.cell(i, 1, t).font = (SUB if t and t[0].isdigit() else NORM if t else NOTE)
    ws3.column_dimensions["A"].width = 115

    out = OUT
    try:
        wb.save(out)
    except PermissionError:
        out = OUT.with_name(OUT.stem + "_v2.xlsx"); wb.save(out)
    print(f"[*] WROTE {out}")


def main():
    itr = sorted(intraday_trades(), key=lambda x: x["entry"])
    ir = np.array([t["net"] for t in itr])
    iy = (date.fromisoformat(itr[-1]["entry"]) - date.fromisoformat(itr[0]["entry"])).days / 365.25
    ds = P.load(); ptr = sorted(positional_trades(ds), key=lambda x: x["entry"])
    pr = np.array([t["net"] / 3.0 for t in ptr])
    py = (date.fromisoformat(ptr[-1]["entry"]) - date.fromisoformat(ptr[0]["entry"])).days / 365.25
    idates = [t["entry"] for t in itr]; pdates = [t["entry"] for t in ptr]
    intr = analyze(ir, idates, iy); posi = analyze(pr, pdates, py)
    for nm, a in [("INTRADAY", intr), ("POSITIONAL", posi)]:
        print(f"\n=== {nm} ===")
        print(f"  observed: total {a['obs_total']:+.0f}% | annual {a['obs_annual']:+.0f}% | maxDD {a['obs_dd']:.0f}%")
        print(f"  bootstrap total 5/50/95: {a['boot_tot_pctiles'][5]:+.0f}/{a['boot_tot_pctiles'][50]:+.0f}/{a['boot_tot_pctiles'][95]:+.0f}%"
              f" | P(total<=0) {a['p_total_le0']:.2f}% | annual 95%CI {a['annual_ci']}")
        print(f"  reshuffle DD 5/50/95: {a['shuf_dd_pctiles'][5]:.0f}/{a['shuf_dd_pctiles'][50]:.0f}/{a['shuf_dd_pctiles'][95]:.0f}%"
              f" | observed DD {a['obs_dd']:.0f}% = {a['obs_dd_pctile']:.0f}th pct | P(losing month) {a['p_lose_month']:.0f}%")
        ov, _ = verdict(a, nm); print("  VERDICT:", ov)
    write(intr, posi)


if __name__ == "__main__":
    import sys
    sys.path.insert(0, r"C:\Users\SPS\Desktop\Kanida.ai Terminal Quant Intelligence Engine\scripts")
    main()
