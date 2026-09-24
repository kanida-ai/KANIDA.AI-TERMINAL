"""Capital deployment check — recommended config, deploy Rs5L EVERY trading day,
SIMPLE (non-compounded) returns. Base stays 5L each day; daily P&L = ret% x 5L
accumulates. Regime gate (self trail-10 > 0): on skip days nothing is deployed
(cash, 0 P&L). Monthly table shows Rs made on wins vs lost on losses + running
capital. GROSS of costs. Writes a new xlsx.
"""
import sys
sys.path.insert(0, r"C:\Users\SPS\Desktop\Kanida.ai Terminal Quant Intelligence Engine\scripts")
sys.path.insert(0, r"C:\Users\SPS\Desktop\Kanida.ai Terminal Quant Intelligence Engine\backend")
from collections import defaultdict
import numpy as np
from falcon_regime_filter import ROOT, strategy_returns, selfreg

DEPLOY = 500000.0
MONTHS = ["", "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def main():
    dates, rets = strategy_returns()          # optimized config daily return %
    gate = selfreg(dates, rets, 10)           # trailing 10-day self return (walk-forward)
    # per-day P&L in Rs (traded only)
    perday = []
    for d in dates:
        traded = gate[d] > 0
        ret = rets[d]
        pnl = (ret / 100.0 * DEPLOY) if traded else 0.0
        perday.append((d, traded, ret, pnl))

    by = defaultdict(list)
    for d, tr, ret, pnl in perday:
        by[d[:7]].append((tr, ret, pnl))

    rows = []
    cum = 0.0
    for ym in sorted(by):
        items = by[ym]
        traded = [x for x in items if x[0]]
        wins = [x for x in traded if x[1] > 0]
        losses = [x for x in traded if x[1] < 0]
        win_pnl = sum(x[2] for x in wins)
        loss_pnl = sum(x[2] for x in losses)
        month_pnl = sum(x[2] for x in traded)
        cum += month_pnl
        rows.append(dict(
            year=int(ym[:4]), month=int(ym[5:7]), traded=len(traded), skip=len(items) - len(traded),
            win=len(wins), loss=len(losses),
            winp=round(len(wins) / len(traded) * 100, 1) if traded else 0.0,
            avgwin=round(win_pnl / len(wins), 0) if wins else 0.0,
            avgloss=round(loss_pnl / len(losses), 0) if losses else 0.0,
            winpnl=round(win_pnl, 0), losspnl=round(loss_pnl, 0),
            monthpnl=round(month_pnl, 0), capital=round(DEPLOY + cum, 0)))

    total_pnl = cum
    print(f"[*] {len(dates)} days | deploy Rs{DEPLOY:,.0f}/day (simple, gated) | traded {sum(1 for _,t,_,_ in perday if t)}")
    print(f"[*] START capital Rs{DEPLOY:,.0f} -> END Rs{DEPLOY+total_pnl:,.0f}  (total P&L +Rs{total_pnl:,.0f})")

    # ---- xlsx ----
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    F = "Calibri"
    TITLE = Font(name=F, bold=True, size=13, color="1F6F8B"); NOTE = Font(name=F, size=9, italic=True, color="666666")
    H = Font(name=F, bold=True, size=9, color="FFFFFF"); HDR = PatternFill("solid", start_color="1F6F8B")
    NORM = Font(name=F, size=10); BOLD = Font(name=F, bold=True, size=10)
    GRN = Font(name=F, size=10, color="1E7E34"); RED = Font(name=F, size=10, color="B00020")
    ALT = PatternFill("solid", start_color="F5FAFB"); TOT = PatternFill("solid", start_color="FFF4CE")
    thin = Side(style="thin", color="DDDDDD"); BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
    Cc = Alignment(horizontal="center"); Rr = Alignment(horizontal="right")
    RS = '#,##0;[Red]-#,##0'
    wb = Workbook(); ws = wb.active; ws.title = "Capital Deployment"
    ws.cell(1, 1, "Capital Deployment — Rs5,00,000/day, SIMPLE (non-compounded)").font = TITLE
    ws.cell(2, 1, "Recommended config: Top-5 cash, arm 2.5 / floor 1 / giveback 1.5 / stop 3, square-off 15:29, regime gate (self trail-10>0). "
                  "Base stays Rs5L each day; daily P&L = day return% x 5L. Skipped (regime-off) days deploy nothing. GROSS of costs.").font = NOTE
    heads = ["Year", "Month", "Traded", "Win", "Loss", "Win%", "Avg Win Rs", "Avg Loss Rs",
             "Won Rs (wins)", "Lost Rs (losses)", "Month P&L Rs", "Capital Rs (5L+cum)"]
    for c, h in enumerate(heads, 1):
        cell = ws.cell(4, c, h); cell.font = H; cell.fill = HDR; cell.border = BORD; cell.alignment = Cc
    r = 5

    def wr(vals, bold=False, fill=None):
        nonlocal r
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v); cell.border = BORD; cell.font = BOLD if bold else NORM
            cell.alignment = Cc if c <= 2 else Rr
            if c in (7, 8, 9, 10, 11, 12): cell.number_format = RS
            if c == 6: cell.number_format = '0.0"%"'
            if fill: cell.fill = fill
            if not bold and c == 11 and isinstance(v, (int, float)):
                cell.font = GRN if v >= 0 else RED
        r += 1

    alt = False
    yr_cum = defaultdict(lambda: [0, 0, 0, 0, 0.0, 0.0, 0.0])  # traded,win,loss,skip,winpnl,losspnl,monthpnl
    for row in rows:
        wr([row["year"], MONTHS[row["month"]], row["traded"], row["win"], row["loss"], row["winp"],
            row["avgwin"], row["avgloss"], row["winpnl"], row["losspnl"], row["monthpnl"], row["capital"]],
           fill=(ALT if alt else None))
        alt = not alt
        yc = yr_cum[row["year"]]
        yc[0] += row["traded"]; yc[1] += row["win"]; yc[2] += row["loss"]; yc[3] += row["skip"]
        yc[4] += row["winpnl"]; yc[5] += row["losspnl"]; yc[6] += row["monthpnl"]
    # yearly rows
    cum = 0.0
    for y in sorted(yr_cum):
        yc = yr_cum[y]; cum += yc[6]
        winp = round(yc[1] / yc[0] * 100, 1) if yc[0] else 0.0
        wr([y, "YEAR", yc[0], yc[1], yc[2], winp, round(yc[4] / yc[1], 0) if yc[1] else 0,
            round(yc[5] / yc[2], 0) if yc[2] else 0, round(yc[4], 0), round(yc[5], 0),
            round(yc[6], 0), round(DEPLOY + cum, 0)], bold=True, fill=TOT)
    # grand total
    tw = sum(yc[1] for yc in yr_cum.values()); tl = sum(yc[2] for yc in yr_cum.values())
    ttr = sum(yc[0] for yc in yr_cum.values()); twp = sum(yc[4] for yc in yr_cum.values()); tlp = sum(yc[5] for yc in yr_cum.values())
    wr(["ALL", "TOTAL", ttr, tw, tl, round(tw / ttr * 100, 1), round(twp / tw, 0) if tw else 0,
        round(tlp / tl, 0) if tl else 0, round(twp, 0), round(tlp, 0), round(total_pnl, 0),
        round(DEPLOY + total_pnl, 0)], bold=True, fill=TOT)
    for i, w in enumerate([6, 7, 7, 6, 6, 7, 12, 12, 14, 15, 13, 16], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C5"
    out = ROOT / "docs" / "ops" / "AUTOTRADE_CAPITAL_DEPLOYMENT.xlsx"
    wb.save(out)
    print(f"[*] WROTE {out}")


if __name__ == "__main__":
    main()
