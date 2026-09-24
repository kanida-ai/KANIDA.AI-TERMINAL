# -*- coding: utf-8 -*-
"""Generate the detailed methodology + design document for the stock-miner (STOCK_MINER_METHODOLOGY.docx)."""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
DL = os.path.join(os.path.expanduser("~"), "Downloads")
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
MINT = RGBColor(0x0B, 0x7A, 0x5A)


def h(text, lvl=1):
    p = doc.add_heading(text, level=lvl)
    for r in p.runs: r.font.color.rgb = MINT
    return p


def para(text, bold=False, italic=False):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold; r.italic = italic; return p


def bullet(text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_lead:
        r = p.add_run(bold_lead + " "); r.bold = True
    p.add_run(text); return p


def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Light Grid Accent 1"
    for i, hd in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = hd
        for r in c.paragraphs[0].runs: r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row): cells[i].text = str(v)
    doc.add_paragraph()
    return t


# ---------------- TITLE ----------------
title = doc.add_heading("Kanida.AI — Per-Stock Pattern Mining Engine", level=0)
for r in title.runs: r.font.color.rgb = MINT
sub = para("Methodology, Design Principles, Results, Limitations & Roadmap", italic=True)
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
para("Module: stock_miner/   |   Universe: Nifty-500 (491 mined, 475 with signals)   |   "
     "Period: 29-Jun-2020 to 27-Jul-2026   |   Prepared: 28-Jul-2026 (IST)").alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph()

# ---------------- 1. EXECUTIVE SUMMARY ----------------
h("1. Executive Summary — What Was Built", 1)
para("A clean-room, per-stock statistical pattern engine that mines each Nifty-500 stock independently and "
     "produces a daily directional signal (Buy / Strong Buy / Sell / Strong Sell / Neutral). It is deliberately "
     "NOT the Falcon engine — it is a separate, isolated system with its own folder, database and feature library, "
     "built to test one question: can per-stock daily patterns produce a tradeable, leak-free edge?")
para("Headline result (validated, leak-free): a genuine and regime-stable SHORT-side edge. The Strong Sell tier "
     "holds 57–59% next-day accuracy in every year 2020–2026. The long side (Buy/Strong Buy ~47%) has no edge — "
     "next-day open-to-close drifts down on the names the trees flag as bullish. At Rs 30,000 per signal, net of "
     "0.10% cost, the short book makes money every year but the edge is thin per-trade and decaying.", bold=False)
para("What this document covers: the exact methodology, the walk-forward design, why it is leak-free, every "
     "indicator used and how, what the system does and does not answer, the intellectual property, and a candid "
     "roadmap on whether 70% win-rate and 1%/day at 1x are reachable.", italic=True)

# ---------------- 2. DESIGN PRINCIPLES ----------------
h("2. Design Principles", 1)
bullet("each stock is mined on its own history and its own patterns — no pooled/global model (a pooled model just "
       "learns market beta). This is what makes it a per-stock engine.", "Per-stock, not pooled.")
bullet("the model may only ever see data strictly before the week it is predicting. Enforced structurally, not by "
       "hope (see Section 4).", "Point-in-time / leak-free by construction.")
bullet("the model retrains every ISO-week on all prior data — it never freezes and adapts as the stock's behaviour "
       "changes.", "Expanding walk-forward, re-mined weekly.")
bullet("a signal is promoted to a tradeable tier only by its RECENT LIVE (out-of-sample) accuracy, never by "
       "training-set purity (training purity overfits and inverts the tiers).", "Tiers earned live, not in training.")
bullet("Buy/Strong Buy = long, Sell/Strong Sell = short. No score, no ranking — a direct actionable call, as "
       "specified.", "Direction, not a score.")
bullet("every artifact — features, rules, model state, realized outcome — is persisted per stock per week so any "
       "pattern can be audited and fine-tuned.", "Everything stored, everything auditable.")
bullet("if a number looks too good, the code is made to report its own arithmetic; bugs (e.g. double-counting "
       "positions) were caught and fixed rather than shipped.", "Rigor over speed.")

# ---------------- 3. METHODOLOGY ----------------
h("3. Mining Methodology (step by step)", 1)
para("For every stock, independently:")
bullet("Load daily OHLCV from 2018-06 (18 months of warm-up so long indicators like SMA-200 are valid before the "
       "test period begins).")
bullet("Compute 129 point-in-time technical features (Section 6). Drop rows before 2020-01-01 and any row with a "
       "missing feature.")
bullet("Walk forward week by week. For test-week W: training set = every day in weeks strictly before W. Minimum "
       "120 training days required, and both up/down outcomes must be present.")
bullet("Train a Decision Tree (max_depth=3, min_samples_leaf=40, fixed random_state). Depth-3 is a deliberate "
       "regularisation choice — it forces each rule to use only 1–3 indicators, which prevents the 129-feature "
       "over-fit that a deeper tree produced (Strong-tier accuracy collapsed to ~40% at depth 4).")
bullet("Each leaf of the tree becomes a candidate pattern: a plain-English rule (e.g. 'seband_pos > 0.52 => down, "
       "70% train accuracy'), a direction, and its training base-rate.")
bullet("Score the test week's actual days through the tree. For each day: which leaf did it fall in, what direction "
       "does that leaf predict, and — critically — what is the rolling LIVE accuracy of that direction over the last "
       "30 realized out-of-sample outcomes for this stock.")
bullet("Assign the tier: Strong (>=15 live samples AND live accuracy >= 65%), Normal (>=8 samples AND >= 55%), else "
       "Neutral. The direction (up/down) sets Buy-family vs Sell-family.")
bullet("Record the outcome: the trade is judged on T+1 open-to-close (enter next day's open, exit next day's close). "
       "Append that result to the rolling live-accuracy memory and move to the next day/week.")
para("Label definition: ret_oc = (next-day close - next-day open) / next-day open x 100. y = 1 if ret_oc > 0. This "
     "is the single, fixed prediction target throughout.", italic=True)

# ---------------- 4. WALK-FORWARD & LEAK-FREE ----------------
h("4. Walk-Forward Method & Why It Is Leak-Free", 1)
para("Method: EXPANDING-window walk-forward at weekly cadence (not a fixed sliding window, not a single train/test "
     "split). The training window grows every week; the test window is always the next single ISO-week; the model is "
     "rebuilt from scratch each week.")
para("Leak-free is guaranteed by four independent structural facts, each verifiable in the source:", bold=True)
bullet("Training rows are filtered to weeks strictly less than the test week (D[D.wk < W]). The model literally "
       "cannot read the test week or any future week.", "Temporal split.")
bullet("Every one of the 129 features is causal — built only from rolling / exponential / shifted windows that look "
       "backward. Iterative indicators (KAMA, McGinley, SuperTrend, NVI/PVI, Connors streak, TD setup) are computed "
       "forward in time. No feature can see its own future.", "Causal features.")
bullet("The tiering uses only PAST realized outcomes — the live-accuracy memory is appended AFTER each day is scored, "
       "so the day being tiered never sees its own result.", "Causal tiering.")
bullet("The label is next-day open-to-close; the features are as-of today's close. There is a full overnight gap "
       "between the last feature and the entry, and no feature is derived from the label.", "No label overlap.")
para("What leak-free does NOT mean: it does not mean profitable. A leak-free system can still lose — and the long "
     "side here does. Leak-free only means the measured accuracy is honest and would have been achievable in real "
     "time.", italic=True)

# ---------------- 5. WHAT IS STORED ----------------
h("5. What Is Stored (per stock, per week)", 1)
para("Database: stock_miner/stock_miner.db (isolated from all production DBs). Plus one parquet per stock.")
table(["Artifact", "Rows / Size", "Contents"], [
    ["signals", "480,179", "Per trade: date, tier, direction, ret_oc, live-accuracy state, the fired rule, train accuracy"],
    ["patterns", "455,675", "Every weekly leaf-rule: symbol, ISO-week, rule text, direction, train accuracy, n_train"],
    ["weekly_state", "101,844", "Per stock/week model state: train size, up-rate, #leaves, live up/down accuracy, samples seen"],
    ["pattern_weekly", "190,770", "Realized OUT-OF-SAMPLE outcome per stock x week x rule x direction (see Section 7)"],
    ["features/<SYM>.parquet", "475 files, 642 MB", "Full 129-feature point-in-time matrix — exactly what the trees saw each day"],
])
para("Interface: stock_miner/engine.py. 'python engine.py <SYMBOL>' writes a per-stock workbook (Q1/Q2 sheets, "
     "feature-edge, weekly state, signals, rules-by-week, full feature matrix). 'python engine.py --winners' ranks "
     "the best/worst stock-feature edges across the whole universe.")

# ---------------- 6. INDICATORS ----------------
h("6. Indicators Used (129 features) and How", 1)
para("All 129 are derived from DAILY OHLCV and expressed as normalised, stationary quantities (percent distances, "
     "positions within bands 0–1, oscillator levels) so they are comparable across price regimes. They are fed as "
     "the candidate feature set to every weekly tree; the depth-3 tree selects the 1–3 most predictive per stock per "
     "week. Families:")
table(["Family", "Examples (indicators)"], [
    ["Trend / moving averages", "SMA/EMA/WMA/HMA/VWMA/KAMA/McGinley/ALMA distances, GMMA ribbon, Ichimoku, Alligator, Aroon, TRIX, DPO, linreg slope/dist, TSF"],
    ["Volatility / bands", "Bollinger %b + width + walking-band count, Keltner, Donchian, STARC, Acceleration, Std-Err, Fib, Chandelier/ATR-stop, Darvas, SuperTrend, Ulcer, ATR%, hist-vol"],
    ["Momentum / strength", "RSI(14/7/3), Stochastic, StochRSI, Williams %R, MACD + histogram, CCI, ROC, CMO, Awesome, ADX/+DI/-DI, Vortex, Ultimate, TSI, Coppock, Fisher, Mass, Centre-of-Gravity, Connors RSI"],
    ["Volume / flow", "Volume ratios/oscillator/ROC/z, OBV, CMF, A/D + Chaikin osc, Ease-of-Movement, MFI, Force, Elder bull/bear power, PVT, NVI/PVI, volume-facilitation"],
    ["Structure / pivots", "Floor pivots + R1/S1, Camarilla, CPR width, distance to 20/60/120/252-day highs/lows, range position, round-number distance, TD setup count"],
    ["Price action", "Candle body, upper/lower wick, gap, close-location-in-range, day range, higher-high / higher-low counts"],
    ["Weekly point-in-time", "Week-to-date return, week position, week range, days-up this week, price vs long MA"],
])
para("Deliberately excluded (data not available at daily grain): intraday VWAP / anchored VWAP, Volume Profile, "
     "Market Facilitation, Fair-Value-Gaps / BPR, session highs/lows, and Open Interest. This exclusion matters — "
     "see Section 8.", italic=True)

# ---------------- 7. WHAT IT ANSWERS ----------------
h("7. What The System Answers", 1)
para("Because the realized outcome of every pattern is stored per week, the engine directly answers, per stock:")
bullet("Which pattern-stock pairs had POSITIVE next-day average behaviour after filtering — pattern_weekly where "
       "pos_behavior = 1 (avg_next_ret > 0); avg_next_ret_filtered isolates the tiered (post-filter) subset.", "Q1.")
bullet("Which patterns FAILED, WHEN, and how many FALSE POSITIVES — pattern_weekly where pos_behavior = 0; "
       "mined_week is exactly when it failed; n_false_pos counts the wrong-direction days.", "Q2.")
para("Universe scale of these answers: 190,770 pattern-weeks; 100,488 (52.7%) positive-behaviour weeks vs 90,282 "
     "failed weeks; 232,176 false positives logged. Accuracy by tier (out-of-sample, all years):")
table(["Tier", "Signals", "Accuracy", "Read"], [
    ["Strong Sell", "49,349", "57.5%", "The edge — short, stable every year"],
    ["Sell", "112,591", "55.5%", "Weak short edge"],
    ["Buy", "22,784", "46.9%", "No edge"],
    ["Strong Buy", "3,338", "47.1%", "No edge (inverted)"],
])
para("Rupee P&L (Rs 30,000/signal, net of 0.10% round-trip, one position per stock per day):")
table(["Book", "Net P&L", "2024", "2025", "2026*"], [
    ["All shorts", "+Rs 8.9 L", "+7.6 L", "+1.6 L", "+0.5 L"],
    ["Strong Sell only", "+Rs 8.2 L", "+4.5 L", "+1.0 L", "+1.1 L"],
    ["All longs", "-Rs 12.1 L", "-2.6 L", "-2.5 L", "-1.7 L"],
    ["Combined", "-Rs 3.2 L", "+5.1 L", "-0.9 L", "-1.2 L"],
])
para("*2026 partial. Also answered: per-stock consistency (83 stocks short-positive every year = CONSISTENT, 163 "
     "MIXED, 228 FINE-TUNE); the per-day basket shape (avg 178 positions/day, ~22 long / 156 short); and the stable "
     "per-stock feature edges (top: SUZLON/ulcer, IDEA/rsi3+bb_width, NCC/atr_tstop, SRF/body — all short).")

# ---------------- 8. WHAT IT DOES NOT ANSWER ----------------
h("8. What The System Does NOT Answer (limitations)", 1)
bullet("It reads daily bars. It cannot see intraday path, timing, order-flow or the chart-reading behaviour that "
       "the human operator's edge actually comes from. This is the single biggest gap.", "No intraday microstructure.")
bullet("A depth-3 rule tells you WHAT co-occurs with the move, not WHY. There is no causal or fundamental model.", "No causation.")
bullet("The top-accuracy names are penny / ultra-high-beta stocks (IDEA, SUZLON, RPOWER, JPPOWER). The 0.10% cost "
       "assumption is optimistic there — real borrow availability and slippage on intraday shorts are worse, and may "
       "erase the thin edge.", "Shorting frictions not modelled.")
bullet("Short edge fell 7.6 L -> 1.6 L -> 0.5 L across 2024-2026. The engine detects decay but does not yet predict "
       "or adapt to it beyond weekly retraining.", "Decay not solved.")
bullet("No long-side edge exists in this design. Half the tiers are effectively unusable.", "Longs dead.")
bullet("Fixed Rs 30k/position. No Kelly / volatility sizing, no portfolio correlation, no capital constraint, no "
       "exit optimisation beyond same-day close.", "No portfolio/exit engineering.")
bullet("Equity cash / MIS only. No options, futures or OI-based signals.", "No F&O.")
para("Most important honest statement: as it stands the system does NOT achieve 70% win-rate, and does NOT achieve "
     "1%/day at 1x. Best tier is ~57.5% accuracy and ~+0.155% gross per trade (~+0.055% net). Sections 9 addresses "
     "whether those targets are reachable.", bold=True)

# ---------------- 9. ROADMAP ----------------
h("9. Roadmap — Can We Reach 70% WR and 1%/day at 1x?", 1)
para("Candid assessment first: 70% win-rate AND 1%/day at 1x, simultaneously and durably, on DAILY-bar next-day "
     "signals, is almost certainly NOT reachable — the daily open-to-close edge tops out near 57–59%. The credible "
     "paths split into 'raise the win-rate / per-trade edge within this design' and 'change the design'.")
para("A. Improvements WITHIN the daily design (realistic: push WR into low-60s, per-trade edge up):", bold=True)
bullet("Trade only CONSISTENT stocks (83) intersected with POSITIVE feature-edges (the fine-tune loop). Fewer, "
       "higher-conviction shorts.", "Restrict the universe.")
bullet("Require confluence: only fire when 2+ independent positive feature-edges agree for that stock and the tier "
       "is Strong. Higher WR, far fewer trades.", "Confluence gate.")
bullet("Replace the fixed same-day close exit with an MFE/MAE-driven trailing exit (the data to calibrate this is "
       "already stored). This targets the +1% per-trade goal directly by letting winners run intraday-to-multiday.", "Exit engineering.")
bullet("A market-regime filter (index breadth / volatility / advance-decline) to sit out the weeks that produced the "
       "2023 and 2025 drawdowns.", "Regime switch.")
bullet("Swap the single depth-3 tree for a small gradient-boosted ensemble and fire only above a high probability "
       "threshold — trades less, but at higher realized accuracy.", "Probability-thresholded ensemble.")
para("B. Design CHANGES (the only credible route to 1%/day and/or 70%):", bold=True)
bullet("Add intraday (1-minute) data and microstructure features. Prior research on this desk shows the operator's "
       "true edge is intraday path-shape, which daily bars cannot reconstruct. This is the highest-leverage change.", "Go intraday.")
bullet("Add the excluded data — Open Interest, VWAP, volume profile, delivery %, sector/peer confirmation, margin/VaR "
       "changes — as fresh feature families.", "New data families.")
bullet("Use leverage honestly: the short edge is intraday (T+1 open-to-close), so it is 5x-able on MIS. A +0.055% "
       "net/trade edge at 5x, high-conviction and selective, is where the 1%-class daily return realistically lives "
       "— NOT at 1x.", "Leverage the intraday edge.")
para("Recommendation: build the fine-tune loop (A) first — it is cheap, uses only stored data, and will tell us the "
     "true ceiling of the daily design — then decide on the intraday build (B), which is the real unlock.", italic=True)

# ---------------- 10. IP ----------------
h("10. Intellectual Property", 1)
para("The defensible IP created here (all isolated in stock_miner/, none of it touching the production Falcon or "
     "AutoTrade execution paths):")
bullet("A per-stock, expanding weekly walk-forward miner with LIVE-accuracy tiering — signals earn their tier from "
       "recent out-of-sample performance, not training purity. This is the novel mechanism.", "Method.")
bullet("A 129-indicator, causal, point-in-time daily feature library (indicators.py) covering seven families, "
       "normalised for cross-regime comparability.", "Feature library.")
bullet("A realized-outcome ledger (pattern_weekly) that records, for every pattern of every stock every week, "
       "whether it worked and how it failed — a reusable substrate for fine-tuning and for training a "
       "winner-vs-loser discriminator.", "Outcome ledger.")
bullet("The empirical finding itself: a leak-free, regime-stable short-side edge on high-beta Indian names, with the "
       "specific stock-feature pairs that carry it.", "The discovery.")
bullet("The per-stock fine-tune engine (engine.py) and the full trade log / journal tooling.", "Tooling.")
para("Ownership & isolation: separate folder, separate database, separate feature code. It reads the shared "
     "read-only market-data DB but writes nothing outside stock_miner/. It is safe to iterate on without any "
     "real-money or production risk.", italic=True)
para("Note on 'IP': this section reads IP as Intellectual Property. If the network/allow-list IP was meant instead, "
     "that is unrelated to this research module — this engine makes no broker connection.", italic=True)

# ---------------- 11. FILES ----------------
h("11. File Map", 1)
table(["File", "Purpose"], [
    ["stock_miner/miner.py", "The parallel weekly walk-forward miner; writes all DB tables + feature parquet"],
    ["stock_miner/indicators.py", "The 129-feature point-in-time indicator library"],
    ["stock_miner/engine.py", "Per-stock fine-tune interface + universe stock-feature ranking"],
    ["stock_miner/report.py", "Month/year accuracy report by tier"],
    ["stock_miner/pnl_pass.py", "Rs 30k long/short P&L, gross & net, by tier and year"],
    ["stock_miner/journal.py", "Stock-level trade log + journal + daily basket view (2024-2026)"],
    ["stock_miner/stock_miner.db", "Isolated database: signals, patterns, weekly_state, pattern_weekly"],
    ["stock_miner/features/", "475 per-stock parquet feature matrices"],
])
para("Deliverables produced in Downloads: STOCK_MINER_MONTHLY.xlsx, STOCK_MINER_PNL.xlsx, STOCK_MINER_JOURNAL.xlsx, "
     "STOCK_MINER_TRADES_2024_2026.csv, PATTERN_STOCK_WINNERS.xlsx, per-stock STOCK_<SYM>.xlsx, and this document.")

out = os.path.join(DL, "STOCK_MINER_METHODOLOGY.docx")
doc.save(out)
print("saved ->", out)
