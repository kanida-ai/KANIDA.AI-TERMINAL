"""FILL the requested monthly grid for EVERY valid (combo x capital x leverage) config, Jan-Jul 2026.
Columns per month: Capital | leverage | Gross ROC% | Rs P&L(gross) | Cost stack Rs | Net P&L Rs | Net ROC% | Max DD | Trading days | Day-win.
Order-type-correct: CNC=1X (long only), MIS=5X (only way to short). Long leg CNC 1X or MIS 5X; short leg MIS 5X fixed.
Combos: Falcon-long, Tail-short, Combo 70/30, Combo 80/20. Capital ladder: 10L, 1cr, 5cr, 10cr.
Writes docs/reports/MONTHLY_CAPACITY_GRID.xlsx (one formatted block per config). Read-only sources; Falcon untouched."""
import os, sqlite3
import numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
ROOT = r"C:\Users\SPS\Desktop\Kanida.ai Terminal Quant Intelligence Engine"
UDB = os.path.join(ROOT, "data", "db", "kanida_universe.db"); RDB = os.path.join(ROOT, "data", "db", "falcon_research.db")

def cost_rs(tb, ts):
    brok = min(20, 0.0003 * tb) + min(20, 0.0003 * ts); stt = 0.00025 * ts; txn = 0.0000297 * (tb + ts)
    stamp = 0.00003 * tb; sebi = 0.000001 * (tb + ts); gst = 0.18 * (brok + txn + sebi)
    return brok + stt + txn + stamp + sebi + gst

uc = sqlite3.connect("file:" + UDB.replace("\\", "/") + "?mode=ro", uri=True)
od = pd.read_sql_query("SELECT symbol, trade_date, open, close, volume FROM ohlc_daily WHERE trade_date BETWEEN '2026-01-01' AND '2026-07-17'", uc); uc.close()
adv = od.assign(tv=od.close * od.volume).groupby("symbol").tv.median().to_dict()
op = od.pivot_table(index="trade_date", columns="symbol", values="open"); cp = od.pivot_table(index="trade_date", columns="symbol", values="close")
cal = sorted(od.trade_date.unique()); nextd = {cal[i]: cal[i + 1] for i in range(len(cal) - 1)}
rc = sqlite3.connect("file:" + RDB.replace("\\", "/") + "?mode=ro", uri=True)
rk = pd.read_sql_query("SELECT signal_date, rank, symbol FROM falcon_full_ranking WHERE signal_date BETWEEN '2025-12-31' AND '2026-07-16'", rc); rc.close()
byday = {sd: g for sd, g in rk.groupby("signal_date")}

def leg(names, ed, sgn, notional):
    names = [s for s in names if ed in op.index and s in op.columns and op.at[ed, s] == op.at[ed, s] and cp.at[ed, s] == cp.at[ed, s] and adv.get(s, 0) > 0]
    if not names or notional <= 0: return 0.0, 0.0, 0.0
    per = notional / len(names); gross = cost = slip = 0.0
    for s in names:
        o = op.at[ed, s]; c = cp.at[ed, s]; qty = per / o
        gross += sgn * (c - o) * qty; cost += cost_rs(o * qty, c * qty)
        part = per / adv[s]; sbps = (2 + 800 * part) / 1e4; slip += 2 * sbps * per
    return gross, cost, slip

def daily(cap, lfrac, llev, sfrac, slev):
    rows = []
    for sd, g in byday.items():
        ed = nextd.get(sd)
        if not ed: continue
        top = list(g[g["rank"] <= 10].symbol); tail = list(g[g["rank"] >= 201].symbol)
        gl = leg(top, ed, 1, cap * lfrac * llev); gs = leg(tail, ed, -1, cap * sfrac * slev)
        G, C, S = gl[0] + gs[0], gl[1] + gs[1], gl[2] + gs[2]
        rows.append((ed, G, C + S, G - C - S))            # cost stack = real costs + slippage
    D = pd.DataFrame(rows, columns=["d", "gross", "cost", "net"]); D["ym"] = D.d.str[:7]
    return D

def monthly(D, cap):
    out = []
    for ym, g in D.groupby("ym"):
        eq = g.net.cumsum(); mdd = (eq.cummax() - eq).max() / cap * 100
        out.append(dict(month=ym, gross_roc=g.gross.sum()/cap*100, gross=g.gross.sum(), cost=g.cost.sum(),
                        net=g.net.sum(), net_roc=g.net.sum()/cap*100, mdd=mdd, days=len(g), daywin=(g.net>0).mean()*100))
    return pd.DataFrame(out)

CAPS = [(1e6, "₹10L"), (1e7, "₹1cr"), (5e7, "₹5cr"), (1e8, "₹10cr")]
# (name, lfrac, sfrac, [(llev, slev, lev_label), ...])   short leg always MIS 5x
CONFIGS = [
    ("Falcon LONG (top-10)", 1.0, 0.0, [(1, 0, "CNC 1X"), (5, 0, "MIS 5X")]),
    ("Tail-short (201-500)", 0.0, 1.0, [(0, 5, "MIS 5X")]),
    ("Combo 70/30",          0.7, 0.3, [(1, 5, "L:CNC1X + S:MIS5X"), (5, 5, "L:MIS5X + S:MIS5X")]),
    ("Combo 80/20",          0.8, 0.2, [(1, 5, "L:CNC1X + S:MIS5X"), (5, 5, "L:MIS5X + S:MIS5X")]),
]
MON = ["2026-01", "2026-02", "2026-03", "2026-04", "2026-05", "2026-06", "2026-07"]
MLBL = {"2026-01": "Jan", "2026-02": "Feb", "2026-03": "Mar", "2026-04": "Apr", "2026-05": "May", "2026-06": "Jun", "2026-07": "Jul"}

# ---- build Excel ----
wb = Workbook(); ws = wb.active; ws.title = "Monthly Capacity Grid"
HDR = ["month", "Capital deployed", "leverage", "Gross ROC (%)", "₹ P&L", "Cost stack", "Net P&L", "Net ROC (%)", "Max DD", "Trading days", "Day-win"]
thin = Side(style="thin", color="BBBBBB"); bd = Border(thin, thin, thin, thin)
hfill = PatternFill("solid", fgColor="1F4E79"); tfill = PatternFill("solid", fgColor="D9E1F2"); cfill = PatternFill("solid", fgColor="F2F2F2")
r = 1
def put(row, vals, bold=False, fill=None, white=False):
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=row, column=j, value=v); c.border = bd; c.alignment = Alignment(horizontal="center")
        if bold: c.font = Font(bold=True, color="FFFFFF" if white else "000000")
        if fill: c.fill = fill

for name, lfrac, sfrac, levs in CONFIGS:
    for llev, slev, llbl in levs:
        for cap, ctag in CAPS:
            D = daily(cap, lfrac, llev, sfrac, slev); m = monthly(D, cap).set_index("month")
            title = f"{name}  —  {ctag}  —  {llbl}"
            tc = ws.cell(row=r, column=1, value=title); tc.font = Font(bold=True, size=12); r += 1
            put(r, HDR, bold=True, fill=hfill, white=True); r += 1
            tot_g = tot_c = tot_n = tot_d = 0; wins = 0
            for ym in MON:
                if ym not in m.index: continue
                x = m.loc[ym]
                put(r, [MLBL[ym], ctag, llbl, round(x.gross_roc, 1), round(x.gross), round(x.cost), round(x.net),
                        round(x.net_roc, 1), f"{x.mdd:.0f}%", int(x.days), f"{x.daywin:.0f}%"])
                tot_g += x.gross; tot_c += x.cost; tot_n += x.net; tot_d += x.days; wins += (x.daywin/100)*x.days
                r += 1
            eqA = D.net.cumsum(); tot_mdd = (eqA.cummax() - eqA).max() / cap * 100
            put(r, ["TOTAL", ctag, llbl, round(tot_g/cap*100, 1), round(tot_g), round(tot_c), round(tot_n),
                    round(tot_n/cap*100, 1), f"{tot_mdd:.0f}%", int(tot_d), f"{wins/tot_d*100:.0f}%"], bold=True, fill=tfill)
            r += 2
for col, w in zip("ABCDEFGHIJK", [7, 15, 20, 13, 16, 15, 16, 12, 9, 13, 9]):
    ws.column_dimensions[col].width = w
outx = os.path.join(ROOT, "docs", "reports", "MONTHLY_CAPACITY_GRID.xlsx"); wb.save(outx)
print("Excel written ->", outx, "\n")

# ---- print the 80/20 blocks inline (both leverages, all capitals) ----
def show(name, lfrac, sfrac, llev, slev, llbl):
    print("=" * 118)
    for cap, ctag in CAPS:
        D = daily(cap, lfrac, llev, sfrac, slev); m = monthly(D, cap).set_index("month")
        print(f"\n{name} — {ctag} — {llbl}")
        print(f"{'month':<6}{'Capital':>9}{'lev':>10}{'GrossROC%':>11}{'₹P&L':>16}{'Cost stack':>15}{'Net P&L':>16}{'NetROC%':>9}{'MaxDD':>7}{'days':>6}{'Daywin':>8}")
        tg = tc = tn = td = w = 0
        for ym in MON:
            if ym not in m.index: continue
            x = m.loc[ym]
            print(f"{MLBL[ym]:<6}{ctag:>9}{llbl:>10}{x.gross_roc:>+11.1f}{x.gross:>+16,.0f}{x.cost:>15,.0f}{x.net:>+16,.0f}{x.net_roc:>+9.1f}{x.mdd:>6.0f}%{int(x.days):>6}{x.daywin:>7.0f}%")
            tg += x.gross; tc += x.cost; tn += x.net; td += x.days; w += (x.daywin/100)*x.days
        eqA = D.net.cumsum(); tmdd = (eqA.cummax()-eqA).max()/cap*100
        print(f"{'TOTAL':<6}{ctag:>9}{llbl:>10}{tg/cap*100:>+11.1f}{tg:>+16,.0f}{tc:>15,.0f}{tn:>+16,.0f}{tn/cap*100:>+9.1f}{tmdd:>6.0f}%{int(td):>6}{w/td*100:>7.0f}%")

print("################  COMBO 80/20  ################")
show("Combo 80/20", 0.8, 0.2, 5, 5, "L5X+S5X")
show("Combo 80/20", 0.8, 0.2, 1, 5, "L1X+S5X")
