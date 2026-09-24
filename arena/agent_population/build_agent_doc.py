"""Deep-dive DOCX for the 6 survivor agents (3 bullish + 3 bearish, 1D reversal patterns).
Pulls EVERYTHING: pattern logic (code + plain-English), role mandate, edge-vs-base across horizons 1/3/5 (1D & 1W),
Rs5L account money stats, best/worst stocks, and an honest verdict. Read-only."""
import importlib.util, inspect, os, glob, textwrap
import numpy as np, pandas as pd
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
ROOT = r"C:\Users\SPS\Desktop\Kanida.ai Terminal Quant Intelligence Engine"
INST = os.path.join(ROOT, "arena", "agent_population", "instances")
CAP = 5e5
BULL = ["RSI oversold reversal", "Spring pattern", "Bullish harami"]
BEAR = ["Reversal after spike bear", "RSI overbought reversal", "Supply zone tap"]
TARGETS = [(s, "bullish") for s in BULL] + [(s, "bearish") for s in BEAR]
SIGN = {"bullish": 1, "bearish": -1}
PLAIN = {
 "RSI oversold reversal": "14-period RSI drops into oversold territory and then turns back up — the agent buys the bounce off an exhausted sell-off.",
 "Spring pattern": "Price briefly stabs below a recent support/range low (a shake-out) and closes back inside — a Wyckoff 'spring'. The agent buys the failed breakdown / reclaim.",
 "Bullish harami": "A small up-candle sits entirely inside the previous large down-candle's body — a compression/indecision reversal after selling. The agent buys the turn.",
 "Reversal after spike bear": "After a sharp upside spike (over-extension), price prints a bearish reversal bar. The agent shorts the fade of the spike.",
 "RSI overbought reversal": "14-period RSI pushes into overbought and then rolls over — the agent shorts the reversal from an over-bought extreme.",
 "Supply zone tap": "Price rallies up into a prior supply/resistance zone and is rejected. The agent shorts the tap of overhead supply.",
}

spec = importlib.util.spec_from_file_location("custom_agent", os.path.join(ROOT, "backend", "_archive", "agents", "custom_agent.py"))
ca = importlib.util.module_from_spec(spec); spec.loader.exec_module(ca)
LBL2FN = {lbl: fn for _, lst in ca.STRATEGY_MAP.items() for lbl, fn in lst}
def src_of(lbl):
    try: return textwrap.dedent(inspect.getsource(LBL2FN[lbl])).strip()
    except Exception: return "(inline lambda — see registry)"
def min_bars(lbl):
    import re
    s = src_of(lbl); g = re.search(r"i\s*<\s*(\d+)", s)
    return int(g.group(1)) if g else "-"

# ---- load instances (only the 6 strategies) ----
led_files = [f for f in glob.glob(os.path.join(INST, "*.csv")) if not f.endswith(".base.csv") and not f.endswith(".acct.csv")]
L, B, A = [], [], []
tset = set(s for s, _ in TARGETS)
for lf in led_files:
    stock = os.path.basename(lf)[:-4]
    d = pd.read_csv(lf); d = d[d.strategy.isin(tset)]; L.append(d)
    B.append(pd.read_csv(lf[:-4] + ".base.csv"))
    af = lf[:-4] + ".acct.csv"
    if os.path.exists(af):
        a = pd.read_csv(af); A.append(a[a.strategy.isin(tset)])
L = pd.concat(L, ignore_index=True); B = pd.concat(B, ignore_index=True); A = pd.concat(A, ignore_index=True)

def ym_i(s): y, m = s.split("-"); return int(y) * 12 + int(m)
def edge_table(strat, bias, tf):
    led = L[(L.strategy == strat) & (L.timeframe == tf)]
    rows = []
    for h in [1, 3, 5]:
        lh = led[led.horizon == h].merge(B[B.horizon == h][["stock", "timeframe", "ym", "base_ret"]],
                                         on=["stock", "timeframe", "ym"], how="left")
        per = lh.groupby("stock").apply(lambda g: pd.Series({
            "n": g.n.sum(), "avg": g.sum_ret.sum() / g.n.sum(),
            "base": (g.base_ret * g.n).sum() / g.n.sum(), "win": g.wins.sum() / g.n.sum() * 100}))
        per["edge"] = per["avg"] - per["base"] * SIGN[bias]
        rows.append((h, len(per), per.n.mean(), per["avg"].median(), per["base"].median(),
                     per["edge"].median(), (per["edge"] > 0).mean() * 100, per["win"].mean()))
    return pd.DataFrame(rows, columns=["h", "stocks", "avg_fires", "agent%", "base%", "edge%", "breadth%", "win%"])

def per_stock_edge(strat, bias):
    lh = L[(L.strategy == strat) & (L.timeframe == "1D") & (L.horizon == 5)].merge(
        B[B.horizon == 5][["stock", "timeframe", "ym", "base_ret"]], on=["stock", "timeframe", "ym"], how="left")
    per = lh.groupby("stock").apply(lambda g: pd.Series({
        "fires": int(g.n.sum()), "agent%": g.sum_ret.sum()/g.n.sum(),
        "base%": (g.base_ret*g.n).sum()/g.n.sum(), "win%": g.wins.sum()/g.n.sum()*100}))
    per["edge%"] = per["agent%"] - per["base%"] * SIGN[bias]
    return per[per.fires >= 20].sort_values("edge%", ascending=False)

def money(strat):
    a = A[(A.strategy == strat) & (A.timeframe == "1D")]
    out = []
    for stk, g in a.groupby("stock"):
        g = g.sort_values("ym"); eq = CAP; peak = CAP; mdd = 0
        for mr in g.month_ret:
            eq *= (1 + mr/100); peak = max(peak, eq); mdd = max(mdd, (peak-eq)/peak)
        span = ym_i(g.ym.iloc[-1]) - ym_i(g.ym.iloc[0]) + 1; tr = int(g.trades.sum())
        if tr >= 20:
            out.append((eq/CAP*100-100, ((eq/CAP)**(12/span)-1)*100 if span >= 6 and eq > 0 else np.nan, mdd*100, tr))
    m = pd.DataFrame(out, columns=["roc", "cagr", "mdd", "tr"])
    return m

# ---- build docx ----
doc = Document()
doc.styles["Normal"].font.name = "Calibri"; doc.styles["Normal"].font.size = Pt(10.5)
h = doc.add_heading("Kanida.AI Arena — Agent Deep-Dive", 0)
doc.add_paragraph("Six survivor patterns from the 127,022-agent population (Addendum G). These are the only "
    "strategies with a POSITIVE, breadth-backed edge over base rate at the daily timeframe — every one a "
    "reversal / mean-reversion pattern. All figures are net of 0.15%/trade cost, measured on 211 F&O stocks, "
    "daily bars 2016–2026, and reported as EDGE-vs-BASE (pattern return minus buy-and-hold of the same stock) so "
    "stock drift is stripped out.").italic = True

def tbl(doc, df, fmts):
    t = doc.add_table(rows=1, cols=len(df.columns)); t.style = "Light Grid Accent 1"
    for j, c in enumerate(df.columns): t.rows[0].cells[j].text = str(c)
    for _, r in df.iterrows():
        cells = t.add_row().cells
        for j, (c, f) in enumerate(zip(df.columns, fmts)):
            cells[j].text = f(r[c])
    return t

for strat, bias in TARGETS:
    doc.add_page_break()
    doc.add_heading(f"{strat}  ({bias})", 1)
    doc.add_paragraph(PLAIN[strat])
    # identity
    doc.add_heading("Identity", 2)
    p = doc.add_paragraph()
    p.add_run("Function: ").bold = True; p.add_run(f"{LBL2FN[strat].__name__}   ")
    p.add_run("Bias: ").bold = True; p.add_run(f"{bias}   ")
    p.add_run("Min-bars N: ").bold = True; p.add_run(f"{min_bars(strat)}   ")
    p.add_run("Agent-ID form: ").bold = True; p.add_run(f"{'BULL' if bias=='bullish' else 'BEAR'}xxx_1D_<STOCK>")
    # logic
    doc.add_heading("Pattern logic (source)", 2)
    code = doc.add_paragraph(src_of(strat)[:1400]); code.style = doc.styles["Normal"]
    for run in code.runs: run.font.name = "Consolas"; run.font.size = Pt(8)
    # role mandate
    doc.add_heading("Role mandate", 2)
    doc.add_paragraph(f"Hypothesis: a '{strat}' on a closed daily bar precedes a "
        f"{'positive (go long)' if bias=='bullish' else 'negative (go short)'} 1–5 bar move.\n"
        f"Universe: each F&O stock (Auto-deployed across all 211).  Timeframe: 1D (1W shown for contrast).\n"
        f"Entry: next-bar open after fire.  Exit: close after 1/3/5 bars (money account uses 5).  "
        f"Direction: {'long' if bias=='bullish' else 'short'}.  Cost: 0.15% round-trip.\n"
        f"Success metric: edge-vs-base, win%, breadth.  Prohibited: no lookahead; no rule changes; no copying the leaderboard winner.")
    # edge tables
    doc.add_heading("Edge-vs-base by holding period — DAILY (1D)", 2)
    e1 = edge_table(strat, bias, "1D")
    tbl(doc, e1.rename(columns={"h": "hold(bars)"}),
        [lambda v: f"{int(v)}", lambda v: f"{int(v)}", lambda v: f"{v:.1f}", lambda v: f"{v:+.3f}",
         lambda v: f"{v:+.3f}", lambda v: f"{v:+.3f}", lambda v: f"{v:.1f}", lambda v: f"{v:.1f}"])
    e1w = edge_table(strat, bias, "1W")
    doc.add_paragraph(f"Weekly (1W) for contrast — 5-bar edge: {e1w[e1w.h==5]['edge%'].iloc[0]:+.3f}%  "
                      f"breadth {e1w[e1w.h==5]['breadth%'].iloc[0]:.0f}%  ({int(e1w[e1w.h==5]['stocks'].iloc[0])} stocks).").italic = True
    # money
    doc.add_heading("₹5 L account (Model A, 1D, 5-bar hold) — money, drift-laden", 2)
    m = money(strat)
    doc.add_paragraph(f"Across {len(m)} stocks (≥20 trades): median ROC {m.roc.median():+.0f}%, median CAGR "
        f"{m.cagr.median():+.2f}%, median max-drawdown {m.mdd.median():.0f}%, avg {m.tr.mean():.0f} trades. "
        f"NOTE: the ₹ figure is inflated by the underlying stock's drift — the EDGE table above is the skill metric.")
    # best/worst stocks
    doc.add_heading("Where it works best / worst (1D, 5-bar, edge-vs-base)", 2)
    ps = per_stock_edge(strat, bias)
    bw = pd.concat([ps.head(5).assign(rank="TOP"), ps.tail(5).assign(rank="BOTTOM")]).reset_index()
    tbl(doc, bw[["rank", "stock", "fires", "agent%", "base%", "edge%", "win%"]],
        [str, str, lambda v: f"{int(v)}", lambda v: f"{v:+.2f}", lambda v: f"{v:+.2f}",
         lambda v: f"{v:+.2f}", lambda v: f"{v:.0f}"])

doc.add_page_break(); doc.add_heading("Honest verdict", 1)
doc.add_paragraph(
    "1) These six are the ONLY breadth-backed positive-edge patterns out of 301 — all reversal / mean-reversion, all daily. "
    "2) The edges are small (+0.1 to +0.6%/trade over base) but robust (positive on 66–82% of stocks). "
    "3) Their ₹5 L account CAGR is often modest or negative because the edge is small and, for the bearish ones, they short in a "
    "decade-long bull market — the ALPHA is real, the absolute money in this regime is not. "
    "4) None of this is out-of-sample yet: the walk-forward calibration (12-mo train / 1-mo OOS) is the real test before any deploy. "
    "This is a research finding, not a live strategy.")
out = os.path.join(ROOT, "docs", "reports", "ARENA_AGENT_DEEPDIVE_6.docx"); doc.save(out)
print("saved ->", out)
print("edge summary (1D h5):")
for strat, bias in TARGETS:
    e = edge_table(strat, bias, "1D"); r = e[e.h == 5].iloc[0]
    print(f"  {strat:<26}{bias:<9} edge {r['edge%']:+.3f}%  breadth {r['breadth%']:.0f}%  stocks {int(r['stocks'])}")
